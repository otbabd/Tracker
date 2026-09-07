"""The consolidator, with a full exercise pushed through it.

Twelve groups' asks and returns — Risk Group's pasted from the worked example
exactly as the central team would paste them — challenged, scored and read back.
Every number the workbook produces is recomputed in Python from the same inputs.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import harness  # noqa: E402  — imported first; it puts capacity/build on the path
from harness import Run
import consolidator as C
import fill as F
import model as M
import refdata as R

SCRATCH = Path("/tmp/claude-0/-home-user-Tracker/"
               "3cac8f7f-9e45-5037-aeaf-db814b62d91c/scratchpad")
TOL = 0.01
CENTS = 0.05


def find_row(ws, label: str, col: int = 1) -> int:
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=col).value == label:
            return r
    raise AssertionError(f"no row labelled {label!r} on {ws.title}")


def prepare(path: Path) -> dict:
    got = F.fill(path, path.parent / "example-for-fill.xlsx")
    result = harness.recalc(path, 900)
    if result.get("status") != "success":
        raise AssertionError(f"recalc: {result}")
    got["recalc"] = result
    return got


def run(prepared: dict) -> Run:
    t = Run("consolidator — full exercise")
    lines, returns = prepared["lines"], prepared["returns"]
    wb = load_workbook(prepared["path"], data_only=True)
    bad = harness.error_cells(wb)
    t.check(not bad, "no cell holds an Excel error", ", ".join(bad[:5]))
    t.equal(prepared["recalc"]["total_errors"], 0, "recalculates with zero errors")

    # ── Central pricing, line by line ──────────────────────────────────────
    ws = wb[C.SH_ASK]
    keys = {}
    for r, line in lines.items():
        want_cat = R.driver_category(line["driver"])
        t.equal(ws[f'{C.K["Category"]}{r}'].value, want_cat, f"line {r} category")
        t.equal(ws[f'{C.K["Rank"]}{r}'].value, R.rank_of(want_cat), f"line {r} rank")
        t.close(TOL, ws[f'{C.K["Rate"]}{r}'].value, M.rate(line["level"]),
                f"line {r} rate")
        t.close(TOL, ws[f'{C.K["One-off"]}{r}'].value, M.one_off(line["level"]),
                f"line {r} one-off")
        t.close(CENTS, ws[f'{C.K["Requested run-rate"]}{r}'].value,
                M.run_rate(line), f"line {r} run-rate cost")
        t.close(CENTS, ws[f'{C.K[f"Requested {R.PLAN_YEAR} cost"]}{r}'].value,
                M.in_year(line), f"line {r} in-year cost")
        t.equal(ws[f'{C.K["Problem"]}{r}'].value, None, f"line {r} has no problem")
        keys[r] = ws[f'{C.K["Sort key"]}{r}'].value
    t.equal(len(set(keys.values())), len(keys), "every sort key is unique")
    t.check(all(k == lines[r]["rank"] * 10**6 + r for r, k in keys.items()),
            "sort keys follow rank then row")

    # ── The challenge rule ─────────────────────────────────────────────────
    chal = wb[C.SH_CHAL]
    approved = {}
    for r, line in lines.items():
        want_total, want_q = M.effective(line)
        t.close(TOL, chal[f"P{r}"].value, want_total, f"challenge {r} effective")
        t.equal(chal[f"J{r}"].value, M.first_quarter(line),
                f"challenge {r} requested start")
        t.equal(chal[f"Q{r}"].value, want_q, f"challenge {r} effective start")
        t.close(CENTS, chal[f"R{r}"].value,
                M.approved_in_year(line, want_total, want_q),
                f"challenge {r} approved cash cost")
        t.close(CENTS, chal[f"S{r}"].value, M.rate(line["level"]) * want_total,
                f"challenge {r} approved run-rate")
        approved[r] = (want_total, want_q)

    # ── Receipts and totals by group ───────────────────────────────────────
    env = wb[C.SH_ENV]
    for group, info in prepared["groups"].items():
        er = info["row"]
        mine = {r: l for r, l in lines.items() if l["group"] == group}
        rets = [x for x in returns.values() if x["group"] == group]
        t.equal(env.cell(row=er, column=13).value, len(mine), f"{group}: rows received")
        t.equal(env.cell(row=er, column=14).value, 0, f"{group}: rows with a problem")
        t.close(TOL, env.cell(row=er, column=15).value,
                sum(M.ask_total(l) for l in mine.values()), f"{group}: requested")
        t.close(CENTS, env.cell(row=er, column=16).value,
                sum(M.run_rate(l) for l in mine.values()),
                f"{group}: requested run-rate")
        t.close(TOL, env.cell(row=er, column=17).value,
                sum(x["exits"] for x in rets), f"{group}: exits")
        t.close(TOL, env.cell(row=er, column=18).value,
                sum(approved[r][0] for r in mine), f"{group}: approved")
        t.close(CENTS, env.cell(row=er, column=19).value,
                sum(M.rate(l["level"]) * approved[r][0] for r, l in mine.items()),
                f"{group}: approved run-rate")

    # ── The bridge balances ────────────────────────────────────────────────
    dash = wb[C.SH_DASH]
    b = prepared_bridge = {
        "opening": dash.cell(row=find_row(
            dash, f"Opening {R.BASE_YEAR} approved establishment"), column=2).value,
        "exits": -dash.cell(row=find_row(
            dash, "less seats the groups are giving up"), column=2).value,
        "approved": dash.cell(row=find_row(
            dash, "plus positions approved"), column=2).value,
        "closing": dash.cell(row=find_row(
            dash, f"Closing {R.PLAN_YEAR} approved establishment"), column=2).value,
    }
    t.close(TOL, b["exits"], sum(x["exits"] for x in returns.values()),
            "exits roll up from the returns")
    t.close(TOL, b["approved"], sum(v[0] for v in approved.values()),
            "approved positions roll up from the challenge")
    t.close(TOL, b["closing"], b["opening"] - b["exits"] + b["approved"],
            "the bridge balances")

    # ── Four scenarios, recomputed from the levers on the sheet ────────────
    scen = wb[C.SH_SCEN]
    eng = wb[C.SH_ENG]
    envelope = env.cell(row=find_row(
        env, f"{R.PLAN_YEAR} cost envelope (SAR '000, run-rate)"), column=2).value
    lever = {name: find_row(scen, name) for name, _, _ in C.LEVERS}
    for k, name in enumerate([s[0] for s in R.SCENARIOS]):
        col = 2 + k
        demand = scen.cell(row=lever["Demand multiplier"], column=col).value
        share = scen.cell(row=lever["Share of envelope released"], column=col).value
        shift = scen.cell(row=lever["Timing shift (quarters)"], column=col).value
        infl = scen.cell(row=lever["Salary inflation"], column=col).value
        prod = scen.cell(row=lever["Productivity multiplier"], column=col).value
        want = M.funded_set(lines, demand, share, shift, infl, prod, envelope)

        mismatch = [r for r, w in want.items()
                    if eng[f"{C.eng_col(k, 3)}{r}"].value != w["funded"]]
        t.check(not mismatch, f"{name}: every line's funding decision matches",
                f"{len(mismatch)} differ, first at row {mismatch[0] if mismatch else ''}")

        funded = [w for w in want.values() if w["funded"]]
        t.equal(scen.cell(row=find_row(scen, "Lines funded"), column=col).value,
                len(funded), f"{name}: lines funded")
        t.close(TOL, scen.cell(row=find_row(scen, "Positions funded"),
                               column=col).value,
                sum(w["total"] for w in funded), f"{name}: positions funded")
        t.close(CENTS, scen.cell(row=find_row(
            scen, "Run-rate cost funded (SAR '000)"), column=col).value,
            sum(w["run_rate"] for w in funded), f"{name}: run-rate funded")
        t.close(CENTS, scen.cell(row=find_row(
            scen, f"{R.PLAN_YEAR} cash cost of those (SAR '000)"), column=col).value,
            sum(w["cash"] for w in funded), f"{name}: cash cost funded")
        t.close(TOL, scen.cell(row=find_row(scen, "Regulatory positions funded"),
                               column=col).value,
                sum(w["total"] for r, w in want.items()
                    if w["funded"] and lines[r]["category"] == "Regulatory"),
                f"{name}: regulatory positions funded")

        # The cut has to fall in ranking order: nothing discretionary survives
        # while something regulatory is dropped.
        worst_funded = max((lines[r]["rank"] for r, w in want.items() if w["funded"]),
                           default=0)
        best_dropped = min((lines[r]["rank"] for r, w in want.items()
                            if not w["funded"]), default=99)
        t.check(worst_funded <= best_dropped,
                f"{name}: the cut respects the ranking order",
                f"funded down to rank {worst_funded}, dropped from {best_dropped}")

    # Deferring starts must not buy headcount: it is a cash lever, not a
    # capacity one. Base and a shifted Base fund exactly the same positions.
    base = M.funded_set(lines, 1.0, 1.0, 0, 0.0, 1.0, envelope)
    shifted = M.funded_set(lines, 1.0, 1.0, 3, 0.0, 1.0, envelope)
    t.check([r for r, w in base.items() if w["funded"]]
            == [r for r, w in shifted.items() if w["funded"]],
            "a timing shift changes the cash, not what gets funded")

    # ── The approved establishment list ────────────────────────────────────
    est = wb[C.SH_EST]
    listed = [r for r in range(C.FIRST_ROW, C.FIRST_ROW + C.EST_ROWS)
              if est.cell(row=r, column=6).value]
    want_lines = [r for r in lines if approved[r][0]]
    t.equal(len(listed), len(want_lines), "every approved line is listed")
    t.close(CENTS, sum(est.cell(row=r, column=10).value or 0 for r in listed),
            sum(approved[r][0] for r in want_lines), "listed positions match")
    t.close(0.5, sum(est.cell(row=r, column=13).value or 0 for r in listed),
            sum(M.rate(lines[r]["level"]) * approved[r][0] for r in want_lines),
            "listed run-rate matches")
    return t


if __name__ == "__main__":
    SCRATCH.mkdir(parents=True, exist_ok=True)
    sys.exit(run(prepare(SCRATCH / "filled.xlsx")).report())
