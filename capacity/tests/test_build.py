"""The three deliverables build, recalculate, and open clean.

A workbook of formulas that has never been recalculated is a guess. This builds
each file from scratch, runs it through LibreOffice, and refuses anything that
comes back with a formula error or an error string left in a cell.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import harness
from harness import Run
import consolidator as C
import example as E
import refdata as R
import template as T

SCRATCH = Path("/tmp/claude-0/-home-user-Tracker/"
               "3cac8f7f-9e45-5037-aeaf-db814b62d91c/scratchpad/build-test")


def run() -> Run:
    t = Run("build — all three workbooks")
    SCRATCH.mkdir(parents=True, exist_ok=True)
    bank = R.Bank()
    t.equal(len(bank.groups), 12, "the sample bank has twelve groups")

    builds = [
        ("template", lambda p: T.build_template(bank.groups[0], p, bank)),
        ("worked example", lambda p: E.build_example(p)),
        ("consolidator", lambda p: C.build_consolidator(p, bank)),
    ]
    for name, build in builds:
        path = SCRATCH / f"{name.replace(' ', '-')}.xlsx"
        build(path)
        t.check(path.exists() and path.stat().st_size > 20_000,
                f"{name}: builds to a real file")
        result = harness.recalc(path, 900)
        t.equal(result.get("status"), "success", f"{name}: recalculates")
        t.equal(result.get("total_errors"), 0, f"{name}: no formula errors")
        t.check(result.get("total_formulas", 0) > 1000,
                f"{name}: is built from formulas, not from pasted numbers",
                f"{result.get('total_formulas')} formulas")

        wb = load_workbook(path, data_only=True)
        bad = harness.error_cells(wb)
        t.check(not bad, f"{name}: no cell holds an error string", ", ".join(bad[:5]))

        wf = load_workbook(path)
        # Protection, dropdowns and the outline all have to survive the save.
        unprotected = [ws.title for ws in wf.worksheets if not ws.protection.sheet]
        if name == "consolidator":
            t.check(not unprotected, f"{name}: every sheet is protected",
                    ", ".join(unprotected))
        else:
            t.equal(unprotected, [T.SH_TREE],
                    f"{name}: only the structure sheet is left unprotected")
            tree = wf[T.SH_TREE]
            t.check(any(d.outline_level for d in tree.row_dimensions.values()),
                    f"{name}: the structure outline survives the save")
        t.check(any(ws.data_validations.dataValidation for ws in wf.worksheets),
                f"{name}: dropdowns survive the save")
        t.check(any(ws.sheet_state == "hidden" for ws in wf.worksheets),
                f"{name}: the reference sheet ships hidden")
        names = set(wf.defined_names)
        for want in ("GradeList", "GradeBasic", "DriverList", "DriverRank",
                     "GosiSaudi", "QuarterList"):
            t.check(want in names, f"{name}: {want} is defined")
    return t


if __name__ == "__main__":
    sys.exit(run().report())
