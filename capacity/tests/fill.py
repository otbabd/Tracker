"""Push a full exercise through the consolidator.

Twelve groups' worth of asks and returns, and a spread of challenge decisions.
Risk Group's bands are filled from the worked example's own sheets, cell for
cell, which is the copy-paste path the central team will actually use; the other
eleven are generated so every band carries load.

Everything written here is returned as plain Python, so the tests can recompute
the workbook's answers without asking the workbook.
"""
from __future__ import annotations

import math
import random
from pathlib import Path

from openpyxl import load_workbook

import harness  # noqa: F401  — puts capacity/build on the path
import consolidator as C
import example as E
import model as M
import refdata as R
import template as T

EXAMPLE_GROUP = E.GROUP

TITLES = ["Analyst", "Senior Analyst", "Officer", "Senior Officer", "Specialist",
          "Senior Specialist", "Manager", "Senior Manager", "Engineer", "Consultant"]
ALTERNATIVES = [
    "Automation considered; the residual work still needs a person.",
    "Vendor cover priced and rejected on cost.",
    "Reviewed against the existing rota; no slack to absorb it.",
    "Considered deferring to 2028; the regulatory date does not move.",
]


def read_example(path: Path, layout: dict) -> tuple[list[dict], list[dict]]:
    """Exactly what a group would copy: the ask block and the return block."""
    wb = load_workbook(path, data_only=True)
    ws = wb[T.SH_ASK]
    A = layout["ask_cols"]
    asks = []
    for r in range(layout["ask"]["first"], layout["ask"]["last"] + 1):
        if not ws[f"{A['Job title']}{r}"].value:
            continue
        asks.append({
            "ref": ws[f"{A['Ref']}{r}"].value,
            "division": ws[f"{A['Division']}{r}"].value,
            "department": ws[f"{A['Department']}{r}"].value,
            "unit": ws[f"{A['Unit']}{r}"].value,
            "sub_unit": ws[f"{A['Sub-Unit']}{r}"].value,
            "title": ws[f"{A['Job title']}{r}"].value,
            "level": ws[f"{A['Career Level']}{r}"].value,
            "family": ws[f"{A['Job Family']}{r}"].value,
            "worker": ws[f"{A['Worker type']}{r}"].value,
            "current": ws[f"{A['Current Capacity']}{r}"].value,
            "total": ws[f"{A['New Asks']}{r}"].value,
            **{q: ws[f"{A[q]}{r}"].value for q in R.QUARTERS},
            "category_sent": ws[f"{A['Category']}{r}"].value,
            "driver": ws[f"{A['Driver']}{r}"].value,
            "alternatives": ws[f"{A['Alternatives considered']}{r}"].value,
        })

    ws = wb[T.SH_POS]
    returns = []
    for r in range(layout["ret"]["first"], layout["ret"]["last"] + 1):
        if ws.cell(row=r, column=1).value in (None, ""):
            continue
        returns.append({
            "division": ws.cell(row=r, column=1).value,
            "current": ws.cell(row=r, column=2).value,
            "exits": ws.cell(row=r, column=3).value,
            "asks": ws.cell(row=r, column=4).value,
            "new": ws.cell(row=r, column=5).value,
            "mandated": ws.cell(row=r, column=6).value,
            "cost": ws.cell(row=r, column=7).value,
        })
    return asks, returns


def generated(bank: R.Bank, group: str, seed: int) -> tuple[list[dict], list[dict]]:
    """A plausible submission for a group whose template nobody filled in."""
    rng = random.Random(seed)
    nodes = bank.descend(bank.group_key(group))
    units = [(k, n) for k, n in nodes if n["level"] >= 2]
    drivers = [d for d, _ in R.ASK_DRIVERS]
    families = bank.job_families or ["Professional"]
    seats = bank.seats(group)

    asks = []
    for i in range(rng.randint(10, 22)):
        key, node = rng.choice(units)
        chain = {}
        k = key
        while k is not None:
            chain[bank.nodes[k]["level"]] = bank.nodes[k]["name"]
            k = bank.nodes[k]["parent"]
        quarters = {q: 0 for q in R.QUARTERS}
        for _ in range(rng.randint(1, 3)):
            quarters[rng.choice(R.QUARTERS)] += 1
        asks.append({
            "ref": i + 1,
            "division": chain.get(1, ""),
            "department": chain.get(2, ""),
            "unit": chain.get(3, chain.get(2, "")),
            "sub_unit": chain.get(4, ""),
            "title": f"{rng.choice(TITLES)} {i + 1}",
            "level": rng.choice(R.CAREER_LEVELS[1:]),
            "family": rng.choice(families),
            "worker": rng.choice(R.WORKFORCE_TYPES[:2]),
            "current": None,
            "total": sum(quarters.values()),
            **quarters,
            "category_sent": None,
            "driver": rng.choice(drivers),
            "alternatives": rng.choice(ALTERNATIVES),
        })

    # The return block, computed the way the group's own sheet computes it —
    # including the label the template gives seats that sit outside a division.
    divisions = [(d, d) for d in bank.divisions(group)]
    if any(not s["division"] for s in seats):
        divisions.append(("Reporting to the group directly", ""))
    returns = []
    for i, (label, div) in enumerate(divisions):
        mine = [s for s in seats if s["division"] == div]
        exits = min(len(mine), 1 + (i + seed) % 3)
        asked = sum(M.ask_total(a) for a in asks if a["division"] == div)
        mandated = 0.0
        cost = 0.0
        for j, seat in enumerate(mine):
            level = E.level_for(seat)
            is_mandated = "No" if j % 3 == 0 else "Yes"
            mandated += 1 if is_mandated == "Yes" else 0
            cost += M.rate_mandated(level, is_mandated)
        returns.append({
            "division": label, "current": float(len(mine)), "exits": float(exits),
            "asks": asked, "new": len(mine) - exits + asked,
            "mandated": mandated, "cost": cost,
        })
    return asks, returns


def decide(i: int, category: str, total: float) -> tuple[str, object, object]:
    """The decision rule the test recomputes from. Deterministic on purpose."""
    if category in ("Regulatory", "Control"):
        return "Approve as asked", None, None
    if i % 11 == 0:
        return "Decline", None, None
    if i % 7 == 0:
        return "Approve fewer", max(1, math.floor(total / 2)), None
    if i % 13 == 0:
        return "Defer to a later quarter", None, "Q4"
    return "Approve as asked", None, None


def fill(path: Path, example_path: Path, bank: R.Bank | None = None) -> dict:
    bank = bank or R.Bank()
    layout = E.build_example(example_path)
    # openpyxl writes formulas without values, so the example has to be
    # recalculated before anything can be copied out of it — which is also the
    # real path: the central team pastes values from a live file.
    result = harness.recalc(example_path, 600)
    if result.get("status") != "success":
        raise AssertionError(f"the worked example did not recalculate: {result}")
    C.build_consolidator(path, bank)

    wb = load_workbook(path)
    ask_ws, ret_ws, chal_ws, env_ws = (wb[C.SH_ASK], wb[C.SH_RET],
                                       wb[C.SH_CHAL], wb[C.SH_ENV])
    env_first = None
    for r in range(1, env_ws.max_row + 1):
        if env_ws.cell(row=r, column=1).value == bank.groups[0]:
            env_first = r
            break
    assert env_first, "could not find the group table on the envelopes sheet"

    lines: dict[int, dict] = {}
    per_group: dict[str, dict] = {}
    all_returns: dict[int, dict] = {}
    for i, group in enumerate(bank.groups):
        if group == EXAMPLE_GROUP:
            asks, returns = read_example(example_path, layout)
        else:
            asks, returns = generated(bank, group, seed=4200 + i)

        bf = C.band_first(i)
        for j, ask in enumerate(asks):
            r = bf + j
            for col, value in [
                ("Ref", ask["ref"]), ("Division", ask["division"]),
                ("Department", ask["department"]), ("Unit", ask["unit"]),
                ("Sub-Unit", ask["sub_unit"]), ("Job title", ask["title"]),
                ("Career Level", ask["level"]), ("Job Family", ask["family"]),
                ("Worker type", ask["worker"]),
                ("Current Capacity", ask["current"]), ("New Asks", ask["total"]),
                ("Q1", ask["Q1"]), ("Q2", ask["Q2"]), ("Q3", ask["Q3"]),
                ("Q4", ask["Q4"]),
                ("Category (as sent)", ask["category_sent"]),
                ("Driver", ask["driver"]),
                ("Alternatives considered", ask["alternatives"]),
            ]:
                ask_ws[f"{C.K[col]}{r}"] = value

            category = R.driver_category(ask["driver"])
            line = dict(ask, group=group, category=category,
                        rank=R.rank_of(category))
            decision, approved, quarter = decide(j, category, M.ask_total(ask))
            chal_ws[f"L{r}"] = decision
            if approved is not None:
                chal_ws[f"M{r}"] = approved
            if quarter is not None:
                chal_ws[f"N{r}"] = quarter
            chal_ws[f"O{r}"] = "Recorded at the challenge meeting"
            line.update(decision=decision, approved=approved,
                        approved_quarter=quarter)
            lines[r] = line

        rf = C.ret_first(i)
        for j, ret in enumerate(returns):
            r = rf + j
            for col, value in [
                ("Division", ret["division"]), ("Current Capacity", ret["current"]),
                ("Exits", ret["exits"]), ("New Asks", ret["asks"]),
                ("New Capacity", ret["new"]), ("Mandated seats", ret["mandated"]),
                ("Run-rate cost", ret["cost"]),
            ]:
                ret_ws[f"{C.RET[col]}{r}"] = value
            all_returns[r] = dict(ret, group=group)

        er = env_first + i
        env_ws.cell(row=er, column=9, value=0.09 + (i % 3) / 100)
        env_ws.cell(row=er, column=10, value="Submitted")
        per_group[group] = dict(row=er, band_first=bf, ret_first=rf,
                                lines=len(asks), returns=len(returns))

    wb.save(path)
    return dict(path=str(path), lines=lines, returns=all_returns,
                groups=per_group, env_first=env_first, first=C.FIRST_ROW,
                last=C.last_row(len(bank.groups)),
                ret_last=C.ret_last(len(bank.groups)), layout=layout)


if __name__ == "__main__":
    import sys

    scratch = Path("/tmp/claude-0/-home-user-Tracker/"
                   "3cac8f7f-9e45-5037-aeaf-db814b62d91c/scratchpad")
    out = Path(sys.argv[1] if len(sys.argv) > 1 else scratch / "filled.xlsx")
    got = fill(out, out.parent / "example-for-fill.xlsx")
    print(f"{got['path']}: {len(got['lines'])} lines, "
          f"{len(got['returns'])} return rows, {len(got['groups'])} groups")
