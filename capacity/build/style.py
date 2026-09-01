"""Shared look for the three workbooks.

The colour convention is the financial-modelling one, so anyone who has opened a
model before can read the sheets without a legend: blue for something you type,
black for something the sheet works out, green for a value pulled from another
sheet, and a yellow fill for the cells you are expected to fill in.
"""
from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"

# Ink
INK = "1A1A1A"
INK_DIM = "4C5563"
INK_MUTE = "7B8694"
BLUE_INPUT = "0000FF"
GREEN_LINK = "008000"
RED_ALERT = "A4262C"
AMBER = "8A5300"
OK_GREEN = "1A6B45"

# Ground — carried over from the org tool so the two read as one family.
NAVY = "24466B"
BAND = "EDF1F6"
RULE = "D6DCE4"
YELLOW = "FFF7CC"
GREY_BG = "F0F2F5"

# Number formats. Percentages are stored as fractions; zeros render as a dash.
MONEY = '#,##0;(#,##0);"-"'
MONEY1 = '#,##0.0;(#,##0.0);"-"'
COUNT = '#,##0;(#,##0);"-"'
FTE = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT0 = '0%;(0%);"-"'

thin = Side(style="thin", color=RULE)
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
UNDER = Border(bottom=thin)


def f(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


TITLE = f(16, True, NAVY)
SUBTITLE = f(10, False, INK_MUTE)
H1 = f(12, True, NAVY)
H2 = f(10, True, INK)
TH = f(9, True, "FFFFFF")
BODY = f(10)
BODY_DIM = f(10, color=INK_DIM)
SMALL = f(9, color=INK_MUTE)
INPUT = f(10, color=BLUE_INPUT)
LINK = f(10, color=GREEN_LINK)
NOTE = f(9, True, color=AMBER)

FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_BAND = PatternFill("solid", fgColor=BAND)
FILL_INPUT = PatternFill("solid", fgColor=YELLOW)
FILL_GREY = PatternFill("solid", fgColor=GREY_BG)

# Sheets ship protected, so an input cell has to say it is open.
UNLOCKED = Protection(locked=False)

WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")
CENTRE = Alignment(horizontal="center", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")


def sheet_title(ws, title: str, subtitle: str = "", width: int = 10) -> int:
    """Title block. Returns the first free row beneath it."""
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws.row_dimensions[1].height = 24
    if subtitle:
        ws["A2"] = subtitle
        ws["A2"].font = SUBTITLE
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(2, width))
    return 4


def header_row(ws, row: int, headers, widths=None, wrap=True) -> None:
    for i, text in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=text)
        c.font = TH
        c.fill = FILL_HEAD
        c.alignment = Alignment(vertical="center", wrap_text=wrap)
        c.border = BOX
    ws.row_dimensions[row].height = 30
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def label_value(ws, row: int, label: str, value=None, fmt=None, font=None,
                note: str = "", col: int = 1):
    ws.cell(row=row, column=col, value=label).font = H2
    c = ws.cell(row=row, column=col + 1, value=value)
    c.font = font or BODY
    if fmt:
        c.number_format = fmt
    if note:
        n = ws.cell(row=row, column=col + 2, value=note)
        n.font = SMALL
        n.alignment = TOP
    return c


def input_cell(c, fmt=None):
    c.font = INPUT
    c.fill = FILL_INPUT
    c.border = BOX
    c.protection = UNLOCKED
    if fmt:
        c.number_format = fmt
    return c


def protect(ws) -> None:
    """Lock everything except the cells marked as inputs. No password: a group
    that genuinely needs to unprotect can, but not by leaning on a key."""
    ws.protection.sheet = True
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False
    ws.protection.sort = False
    ws.protection.autoFilter = False


def band(ws, row: int, cols: int) -> None:
    for i in range(1, cols + 1):
        ws.cell(row=row, column=i).fill = FILL_BAND


def note_line(ws, row: int, text: str, cols: int = 8):
    c = ws.cell(row=row, column=1, value=text)
    c.font = SMALL
    c.alignment = WRAP
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=cols)
    ws.row_dimensions[row].height = 26
    return c


def legend(ws, row: int) -> int:
    """The colour key, so nobody has to guess which cells are theirs."""
    ws.cell(row=row, column=1, value="How to read this workbook").font = H2
    items = [
        ("Type here", INPUT, FILL_INPUT, "Yellow fill, blue text — the cells you fill in."),
        ("Calculated", BODY, None, "Black — worked out by the sheet. Do not overwrite."),
        ("From another sheet", LINK, None, "Green — pulled from elsewhere in this file."),
    ]
    for i, (word, font, fill, why) in enumerate(items, start=1):
        c = ws.cell(row=row + i, column=1, value=word)
        c.font = font
        if fill:
            c.fill = fill
        c.border = BOX
        d = ws.cell(row=row + i, column=2, value=why)
        d.font = SMALL
        ws.merge_cells(start_row=row + i, start_column=2, end_row=row + i, end_column=7)
    return row + len(items) + 2


def print_setup(ws, landscape: bool = True, fit_width: int = 1, fit_height=None,
                title_rows: str = "") -> None:
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = fit_width
    ws.page_setup.fitToHeight = fit_height if fit_height is not None else 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    if title_rows:
        ws.print_title_rows = title_rows
