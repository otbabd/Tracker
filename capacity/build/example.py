"""The worked example: one group's template, filled in the way it should be.

This is the file that goes in the kickoff pack. It is the same template every
group receives — nothing is added to it — with a plausible set of answers typed
into the cells that are theirs, so a group head can see what a good submission
looks like before starting their own.

It also stands in for the one thing the org export cannot supply: a career level
against every seat. Here they are filled in from the sample's job families; in
the real exercise the bank supplies them.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

import refdata as R
import template as T

GROUP = "Risk Group"
DIST = Path(__file__).resolve().parents[1] / "dist"

# The sample export's job families happen to read like a level ladder, which is
# what makes them usable here. Real data will need the bank's own mapping.
FAMILY_TO_LEVEL = {
    "Executive": "Executive",
    "Leadership": "Senior Management",
    "Management": "Management",
    "Professional": "Professional",
}
TITLE_TO_LEVEL = [
    ("head of", "Senior Management"),
    ("chief", "Executive"),
    ("manager", "Management"),
    ("team lead", "Management"),
    ("senior", "Professional"),
    ("specialist", "Professional"),
    ("analyst", "Officer"),
    ("officer", "Officer"),
    ("associate", "Officer"),
    ("assistant", "Support"),
]


def level_for(seat: dict) -> str:
    """What the bank would fill in: the level of the job, not of the grade."""
    low = seat["title"].lower()
    for key, level in TITLE_TO_LEVEL:
        if key in low:
            return level
    return FAMILY_TO_LEVEL.get(seat["family"], "Professional")


def place(bank: R.Bank, group: str, keyword: str) -> dict:
    """Find a real unit by name and fill in the ladder above it.

    An ask has to name a unit that exists, and a division that matches it, or the
    check sheet is right to reject it. Rather than typing four columns by hand
    and getting one wrong, look the placement up.
    """
    key = None
    for k, node in bank.descend(bank.group_key(group)):
        if node["level"] >= 2 and keyword.lower() in node["name"].lower():
            key = k
            break
    if key is None:
        key = next(k for k, n in bank.descend(bank.group_key(group))
                   if n["level"] == 2)
    chain = {}
    while key is not None:
        node = bank.nodes[key]
        chain[node["level"]] = node["name"]
        key = node["parent"]
    return {
        "division": chain.get(1, ""),
        "department": chain.get(2, ""),
        "unit": chain.get(3, chain.get(2, "")),
        "sub_unit": chain.get(4, ""),
    }


# The asks. Written by hand rather than generated: the point of the example is
# that the justifications read like a person wrote them.
#   (unit keyword, job title, career level, worker type, {quarter: number},
#    driver, alternatives considered)
ASKS = [
    ("Risk Modelling", "Senior Model Validation Analyst", "Professional", "Permanent",
     {"Q1": 1}, "SAMA regulatory requirement",
     "Reviewed with Model Risk; the existing two validators cannot cover the "
     "IFRS 9 revalidation and the new scorecards in the same year."),
    ("Risk Modelling", "Model Validation Analyst", "Officer", "Permanent",
     {"Q2": 1}, "New regulation implementation",
     "Considered using the vendor panel; rejected on cost and on the "
     "independence requirement."),
    ("Risk Modelling", "Head of Model Validation", "Senior Management", "Permanent",
     {"Q1": 1}, "SAMA regulatory requirement",
     "Cannot be covered by the existing Head of Model Risk without breaching "
     "the separation SAMA asked for in the 2026 review."),
    ("Corporate Credit Risk", "Credit Risk Analyst", "Officer", "Permanent",
     {"Q2": 1, "Q3": 1}, "Volume growth - existing product",
     "Automation of the pre-screen took out about 15% of the manual work; the "
     "residual volume still needs two more analysts."),
    ("Credit Administration", "Senior Credit Risk Analyst", "Professional", "Permanent",
     {"Q3": 1}, "Volume growth - existing product",
     "Start held to Q3 so the cost lands half-year."),
    ("Operational Risk", "Operational Risk Officer", "Officer", "Permanent",
     {"Q1": 1}, "Backfill - resignation",
     "Same role, same level. The work does not stop when the person leaves."),
    ("Operational Risk", "Operational Risk Officer", "Officer", "Permanent",
     {"Q2": 1}, "Backfill - retirement",
     "Retirement known since 2025; no change to the role."),
    ("Business Continuity", "Resilience Analyst", "Officer", "Permanent",
     {"Q2": 1}, "Control gap closure",
     "The 2026 internal audit finding on out-of-hours cover cannot be closed "
     "with the current five-person rota."),
    ("Business Continuity", "Senior Resilience Specialist", "Professional", "Permanent",
     {"Q3": 1}, "Internal audit finding",
     "Considered extending the outsourced night shift; rejected because the "
     "finding is specifically about decision authority sitting outside the bank."),
    ("IFRS 9", "Risk Reporting Analyst", "Officer", "Permanent",
     {"Q2": 1}, "Insourcing from vendor",
     "Currently an outsourced seat at a higher day rate. Converting is cheaper "
     "from month nine and keeps the reporting logic in-house."),
    ("Risk Analytics", "Risk Data Engineer", "Professional", "Permanent",
     {"Q3": 1}, "Automation enablement",
     "Pays back inside eighteen months on the manual reconciliation it removes."),
    ("Market Risk", "Market Risk Analyst", "Officer", "Permanent",
     {"Q2": 1}, "Risk framework uplift",
     "The revised limit framework needs daily monitoring the current team "
     "cannot absorb on top of the FRTB work."),
    ("Liquidity Risk", "Liquidity Risk Analyst", "Officer", "Permanent",
     {"Q4": 1}, "Regulatory audit finding",
     "Deferred to Q4 deliberately; the framework work has to land first."),
    ("Retail Credit Risk", "Portfolio Monitoring Officer", "Officer", "Insourced",
     {"Q3": 1}, "Volume growth - existing product",
     "Insourced rather than permanent because the volume peak is expected to "
     "unwind once the 2026 vintage runs off."),
]

# Seats the group is giving up. Keyword and job-title fragment, so the example
# does not depend on which row the export happens to put them on.
EXITS = [
    ("Business Continuity", "Analyst"),
    ("Credit Administration", "Associate"),
    ("Market Risk", "Analyst"),
    ("Liquidity Risk", "Officer"),
    ("Retail Credit Risk", "Analyst"),
]


def fill(layout: dict, bank: R.Bank, path: Path) -> Path:
    wb = load_workbook(path)
    C, A = layout["cap_cols"], layout["ask_cols"]

    # ── 2. Current capacity ────────────────────────────────────────────────
    ws = wb[T.SH_CAP]
    seats = bank.seats(GROUP)
    first = layout["cap"]["first"]
    exits_wanted = list(EXITS)
    for i, seat in enumerate(seats):
        r = first + i
        ws[f"{C['Career Level']}{r}"] = level_for(seat)
        # Two seats in three are mandated Saudi. A real group answers this seat
        # by seat; the pattern here just has to be plausible and reproducible.
        ws[f"{C['Nationality Mandate?']}{r}"] = "No" if i % 3 == 0 else "Yes"

        direction = "Hold"
        for j, (unit_key, title_key) in enumerate(exits_wanted):
            if (seat["vacant"] and unit_key.lower() in seat["unit"].lower()
                    and title_key.lower() in seat["title"].lower()):
                direction = R.EXIT_DIRECTION
                exits_wanted.pop(j)
                break
        else:
            if i % 17 == 3:
                direction = "Grow"
            elif i % 29 == 5:
                direction = "Reduce"
        ws[f"{C['Capacity Direction']}{r}"] = direction

    # ── 3. Capacity asks ───────────────────────────────────────────────────
    ws = wb[T.SH_ASK]
    first = layout["ask"]["first"]
    families = bank.job_families
    for i, (unit_key, title, level, worker, quarters, driver, why) in enumerate(ASKS):
        r = first + i
        where = place(bank, GROUP, unit_key)
        ws[f"{A['Division']}{r}"] = where["division"]
        ws[f"{A['Department']}{r}"] = where["department"]
        ws[f"{A['Unit']}{r}"] = where["unit"]
        ws[f"{A['Sub-Unit']}{r}"] = where["sub_unit"] or None
        ws[f"{A['Job title']}{r}"] = title
        ws[f"{A['Career Level']}{r}"] = level
        # The job family is the bank's own catalogue; pick the one that matches
        # the level where the sample's families line up with it.
        ws[f"{A['Job Family']}{r}"] = next(
            (f for f, lv in FAMILY_TO_LEVEL.items() if lv == level),
            families[-1] if families else "Professional")
        ws[f"{A['Worker type']}{r}"] = worker
        for q, n in quarters.items():
            ws[f"{A[q]}{r}"] = n
        ws[f"{A['Driver']}{r}"] = driver
        ws[f"{A['Alternatives considered']}{r}"] = why

    # ── 5. Check & submit ──────────────────────────────────────────────────
    ws = wb[T.SH_CHK]
    chk = layout["chk"]
    ws.cell(row=chk["ready_row"] + 1, column=2,
            value="Head of Risk Portfolio & Planning")
    ws.cell(row=chk["ready_row"] + 2, column=2, value="risk.planning@bank.example")
    ws.cell(row=chk["ready_row"] + 3, column=2, value=f"15 October {R.BASE_YEAR}")
    ws.cell(row=chk["state_row"], column=2, value="Submitted")

    wb.save(path)
    return path


def build_example(path=None) -> dict:
    bank = R.Bank()
    path = Path(path or DIST / f"{R.PLAN_YEAR}-Capacity-Example-RiskGroup.xlsx")
    layout = T.build_template(GROUP, path, bank)
    fill(layout, bank, path)
    layout["path"] = str(path)
    return layout


if __name__ == "__main__":
    print(build_example()["path"])
