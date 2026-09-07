"""The consolidator: twelve returned templates in, one decision pack out.

The central team pastes each group's ask block into its own band and its return
block underneath, prices every line again on the same rate card the groups were
given, ranks the whole book by the agreed order, challenges it line by line, and
reads the result four ways at once. Formulas only — nothing here needs a macro or
a refresh.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import refdata as R
import style as S
import tables as TB
import template as T

DIST = Path(__file__).resolve().parents[1] / "dist"

INTAKE_ROWS = T.ASK_ROWS        # exactly the template's, so a paste cannot overflow
RETURN_ROWS = 8                 # divisions per group, with room to spare
EST_ROWS = 600
HEAD_ROW = 5
FIRST_ROW = HEAD_ROW + 1

SH_READ = "Read me"
SH_ENV = "1. Envelopes"
SH_ASK = "2. Group asks"
SH_RET = "3. Group returns"
SH_CHAL = "4. Challenge"
SH_SCEN = "5. Scenarios"
SH_DASH = "6. Dashboard"
SH_EXEC = "7. Executive summary"
SH_EST = "8. Approved establishment"
SH_PAGES = "9. Group pages"
SH_ENG = "Engine"
SH_REF = "Ref"

QSET = '{"Q1";"Q2";"Q3";"Q4"}'
DECISIONS = ["Approve as asked", "Approve fewer", "Defer to a later quarter", "Decline"]

# The paste target: the template's own ask columns, in the template's own order.
PASTE = ["Ref", "Division", "Department", "Unit", "Sub-Unit", "Job title",
         "Career Level", "Job Family", "Worker type", "Current Capacity",
         "New Asks", "Q1", "Q2", "Q3", "Q4", "Category (as sent)", "Driver",
         "Alternatives considered"]
COMPUTED = ["Category", "Rank", "Rate", "One-off", "Requested run-rate",
            f"Requested {R.PLAN_YEAR} cost", "Sort key", "Problem"]
K = {name: get_column_letter(i + 2) for i, name in enumerate(PASTE)}
K.update({name: get_column_letter(len(PASTE) + 2 + i)
          for i, name in enumerate(COMPUTED)})
K["Group"] = "A"

RETURN_COLS = ["Division", "Current Capacity", "Exits", "New Asks",
               "New Capacity", "Mandated seats", "Run-rate cost"]
RET = {name: get_column_letter(i + 2) for i, name in enumerate(RETURN_COLS)}
RET["Group"] = "A"


def band_first(i: int) -> int:
    return FIRST_ROW + i * INTAKE_ROWS


def ret_first(i: int) -> int:
    return FIRST_ROW + i * RETURN_ROWS


def last_row(n: int) -> int:
    return FIRST_ROW + n * INTAKE_ROWS - 1


def ret_last(n: int) -> int:
    return FIRST_ROW + n * RETURN_ROWS - 1


def _name(wb, sheet: str, name: str, ref: str) -> None:
    wb.defined_names.add(DefinedName(name, attr_text=f"'{sheet}'!{ref}"))


def qidx(cell: str) -> str:
    return f"MATCH({cell},{QSET},0)"


def rate_of(level_cell: str) -> str:
    """Full-year loaded cost of one seat at the level in `level_cell`.

    Per row, so a plain INDEX/MATCH does the job; a blank level yields #N/A,
    which IFERROR turns into nothing rather than into a wrong number.
    """
    m = f"MATCH({level_cell},LevelList,0)"
    return (f"IFERROR(INDEX(LevelBasic,{m})+INDEX(LevelHousing,{m})"
            f"+INDEX(LevelTransport,{m})"
            f"+INDEX(LevelBasic,{m})*INDEX(LevelBonus,{m})"
            f"+(INDEX(LevelBasic,{m})+INDEX(LevelHousing,{m}))*GosiSaudi,0)")


# ── Ref ─────────────────────────────────────────────────────────────────────
def build_ref(wb: Workbook, bank: R.Bank) -> None:
    ws = wb.create_sheet(SH_REF)
    ws.sheet_state = "hidden"
    row = S.sheet_title(ws, "Reference data",
                        "The same tables the groups were given. Do not edit mid-exercise.")
    row = TB.write_settings(wb, ws, SH_REF, row)
    row = TB.write_rate_card(wb, ws, SH_REF, row)
    row = TB.write_drivers(wb, ws, SH_REF, row)
    TB.write_lists(wb, ws, SH_REF, row,
                   extra=[("CareerLevels", R.CAREER_LEVELS),
                          ("JobFamilies", bank.job_families or ["Professional"]),
                          ("DecisionList", DECISIONS),
                          ("GroupList", bank.groups)])


# ── Read me ─────────────────────────────────────────────────────────────────
def build_readme(wb: Workbook, bank: R.Bank) -> None:
    ws = wb.create_sheet(SH_READ)
    row = S.sheet_title(ws, f"{R.PLAN_YEAR} Capacity Exercise — consolidator",
                        "For the central team. The order to work in, and what each "
                        "sheet is for.")
    ws.column_dimensions["A"].width = 28
    for c in "BCDEFGH":
        ws.column_dimensions[c].width = 14

    steps = [
        (SH_ENV,
         "Set the bank's ceilings and each group's share of them, and mark each "
         "return as it arrives. Everything else on this sheet fills itself."),
        (SH_ASK,
         f"One band of {INTAKE_ROWS} rows per group. Open the group's file, copy "
         f"A6:R205 from its {T.SH_ASK} sheet, and paste-special values into that "
         "group's band starting at column B. The band is exactly the size of the "
         "template, so a paste cannot spill into the next group. Everything from "
         "the Category column rightwards is recomputed here on the central rate card."),
        (SH_RET,
         f"The small block at the foot of each group's {T.SH_POS} sheet: what it "
         "has today, what it is giving up, and how many of its seats have to be "
         "Saudi. Copy it in the same way. This is where the bridge gets its exits."),
        (SH_CHAL,
         "Every line, requested against approved. Set a decision on each line; only "
         "'Approve fewer' needs a number typed. Approved — not requested — is what "
         "rolls forward into every output."),
        (SH_SCEN,
         "Four cases, six levers each, computed side by side over the whole book. "
         "The ranking rule funds regulatory first and BAU last, so a tighter "
         "envelope shows you exactly which lines fall out."),
        (SH_DASH,
         f"The four analyses: the {R.BASE_YEAR}→{R.PLAN_YEAR} establishment bridge, "
         "the group league table against envelope, demand by driver asked against "
         "approved, and how the cost phases across the year."),
        (SH_EXEC, "One page for the committee."),
        (SH_EST, "The approved lines as a flat list, ready for Finance and recruitment."),
        (SH_PAGES, "One printable page per group, for sending the answer back."),
    ]
    ws.cell(row=row, column=1, value="The order to work in").font = S.H1
    row += 1
    for name, why in steps:
        ws.cell(row=row, column=1, value=name).font = S.H2
        ws.cell(row=row, column=1).alignment = S.TOP
        c = ws.cell(row=row, column=2, value=why)
        c.font = S.BODY_DIM
        c.alignment = S.WRAP
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=8)
        ws.row_dimensions[row].height = 48
        row += 1
    row += 1
    row = S.legend(ws, row)
    ws.cell(row=row, column=1, value="Three things to hold on to").font = S.H1
    row += 1
    for text in [
        "Every line is priced here on the central rate card, not on whatever the "
        "group typed. If a group's number and yours differ, it is the rate card that "
        "has moved, and the card on the hidden Ref sheet is the only place to change it.",
        "A line with no career level cannot be priced and costs nothing. That is "
        "deliberate — a missing level shows up as a dash rather than as a plausible "
        "wrong number — and the group's own check sheet blocks submission until every "
        "seat has one.",
        "The ranking rule is a starting point, not a decision. It funds regulatory "
        "work first and business as usual last; the Challenge sheet is where a human "
        "overrides it, and the reason column is what you will be asked for afterwards.",
    ]:
        S.note_line(ws, row, text)
        row += 1
    S.print_setup(ws, landscape=False)


# ── 1. Envelopes ────────────────────────────────────────────────────────────
def build_envelopes(wb: Workbook, bank: R.Bank) -> dict:
    ws = wb.create_sheet(SH_ENV)
    row = S.sheet_title(
        ws, "Envelopes and receipts",
        "What the bank can afford, how it is split, and what has come back so far.")

    tot_cost = sum(bank.group_totals(g)["cost"] for g in bank.groups)
    tot_appr = sum(bank.group_totals(g)["approved"] for g in bank.groups)
    tot_saudi = sum(bank.group_totals(g)["saudi"] for g in bank.groups)
    tot_filled = sum(bank.group_totals(g)["filled"] for g in bank.groups)

    ws.cell(row=row, column=1, value="Bank ceiling").font = S.H1
    row += 1
    for name, label, value, fmt, why in [
        ("BankCostEnvelope", f"{R.PLAN_YEAR} cost envelope (SAR '000, run-rate)",
         round(tot_cost * 0.06, 0), S.MONEY,
         "Illustrative: six per cent of the current pay bill."),
        ("BankHeadEnvelope", "Net establishment growth allowed",
         round(tot_appr * 0.05), S.COUNT,
         "Illustrative: five per cent of the current establishment."),
        ("BankSaudiTarget", "Mandated-seat target",
         round(tot_saudi / tot_filled + 0.02, 2), S.PCT,
         "Illustrative: two points above where the bank stands today."),
    ]:
        S.input_cell(S.label_value(ws, row, label, value, fmt, note=why), fmt)
        _name(wb, SH_ENV, name, f"$B${row}")
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="By group").font = S.H1
    ws.cell(row=row, column=10,
            value="Blue cells are yours: the allocation, the attrition rate and the "
                  "state as each return arrives.").font = S.NOTE
    row += 1
    heads = ["Group", f"{R.BASE_YEAR} approved", "Filled", "Vacant",
             f"{R.BASE_YEAR} cost", "Cost envelope", "Headcount envelope",
             "Mandated target", "Attrition rate", "State",
             "Paste asks into", "Paste return into", "Rows received",
             "Rows with a problem", "Requested asks", "Requested run-rate",
             "Exits", "Approved asks", "Approved run-rate", "Against envelope"]
    S.header_row(ws, row, heads,
                 [30, 12, 9, 9, 13, 13, 13, 12, 11, 12, 15, 15, 11, 12,
                  12, 14, 9, 12, 14, 16])
    head = row
    first = row + 1
    n = len(bank.groups)
    ask_g = _rng(SH_ASK, K["Group"], FIRST_ROW, last_row(n))
    ask_t = _rng(SH_ASK, K["Job title"], FIRST_ROW, last_row(n))
    ask_p = _rng(SH_ASK, K["Problem"], FIRST_ROW, last_row(n))
    ask_n = _rng(SH_ASK, K["New Asks"], FIRST_ROW, last_row(n))
    ask_rr = _rng(SH_ASK, K["Requested run-rate"], FIRST_ROW, last_row(n))
    ret_g = _rng(SH_RET, RET["Group"], FIRST_ROW, ret_last(n))
    ret_x = _rng(SH_RET, RET["Exits"], FIRST_ROW, ret_last(n))
    ch_g = _rng(SH_CHAL, "A", FIRST_ROW, last_row(n))
    ch_fte = _rng(SH_CHAL, "P", FIRST_ROW, last_row(n))
    ch_rr = _rng(SH_CHAL, "S", FIRST_ROW, last_row(n))

    for i, g in enumerate(bank.groups):
        r = first + i
        t = bank.group_totals(g)
        q = g.replace('"', '""')
        ws.cell(row=r, column=1, value=g).font = S.H2
        for col, value, fmt in [(2, t["approved"], S.COUNT), (3, t["filled"], S.COUNT),
                                (4, t["vacant"], S.COUNT),
                                (5, round(t["cost"], 0), S.MONEY)]:
            ws.cell(row=r, column=col, value=value).number_format = fmt
        base_saudi = t["saudi"] / t["filled"] if t["filled"] else 0
        for col, value, fmt in [
                (6, round(t["cost"] * 0.06, 0), S.MONEY),
                (7, max(5, round(t["approved"] * 0.05)), S.COUNT),
                (8, round(base_saudi + 0.02, 2), S.PCT),
                (9, R.DEFAULT_ATTRITION, S.PCT), (10, "Draft", None)]:
            S.input_cell(ws.cell(row=r, column=col, value=value), fmt)
        bf, rf = band_first(i), ret_first(i)
        ws.cell(row=r, column=11,
                value=f'{K["Ref"]}{bf}:{K["Alternatives considered"]}'
                      f"{bf + INTAKE_ROWS - 1}").font = S.SMALL
        ws.cell(row=r, column=12,
                value=f'{RET["Division"]}{rf}:{RET["Run-rate cost"]}'
                      f"{rf + RETURN_ROWS - 1}").font = S.SMALL
        for col, formula, fmt in [
            (13, f'=COUNTIFS({ask_g},"{q}",{ask_t},"<>")', S.COUNT),
            (14, f'=COUNTIFS({ask_g},"{q}",{ask_p},"?*")', S.COUNT),
            (15, f'=SUMIFS({ask_n},{ask_g},"{q}")', S.FTE),
            (16, f'=SUMIFS({ask_rr},{ask_g},"{q}")', S.MONEY),
            (17, f'=SUMIFS({ret_x},{ret_g},"{q}")', S.FTE),
            (18, f'=SUMIFS({ch_fte},{ch_g},"{q}")', S.FTE),
            (19, f'=SUMIFS({ch_rr},{ch_g},"{q}")', S.MONEY),
        ]:
            ws.cell(row=r, column=col, value=formula).number_format = fmt
        ws.cell(row=r, column=20,
                value=f'=IF($S{r}<=$F{r},"Within","Over by "&TEXT($S{r}-$F{r},"#,##0"))')
        for col in range(1, 21):
            ws.cell(row=r, column=col).border = S.BOX
    last = first + n - 1

    r = last + 1
    ws.cell(row=r, column=1, value="Bank total").font = S.f(10, True, S.NAVY)
    for col in [2, 3, 4, 5, 6, 7, 13, 14, 15, 16, 17, 18, 19]:
        letter = get_column_letter(col)
        c = ws.cell(row=r, column=col, value=f"=SUM({letter}{first}:{letter}{last})")
        c.number_format = (S.MONEY if col in (5, 6, 16, 19)
                           else S.FTE if col in (15, 17, 18) else S.COUNT)
        c.font = S.f(10, True)
    ws.cell(row=r, column=20,
            value=f'=IF($S{r}<=BankCostEnvelope,"Within the bank envelope",'
                  f'"Over the bank envelope by "&TEXT($S{r}-BankCostEnvelope,"#,##0"))'
            ).font = S.f(10, True)
    total_row = r
    S.band(ws, r, 20)

    r += 2
    ws.cell(row=r, column=1, value="Allocated against the ceiling").font = S.H2
    ws.cell(row=r, column=2,
            value=f'=IF(SUM($F${first}:$F${last})<=BankCostEnvelope,'
                  f'"Group cost envelopes fit inside the bank envelope",'
                  f'"Group cost envelopes exceed the bank envelope by "'
                  f'&TEXT(SUM($F${first}:$F${last})-BankCostEnvelope,"#,##0"))')
    alloc_row = r

    dv = DataValidation(type="list", formula1='"Draft,Submitted,Challenged,Approved"',
                        allow_blank=True, showErrorMessage=True)
    ws.add_data_validation(dv)
    dv.add(f"J{first}:J{last}")
    ws.conditional_formatting.add(
        f"T{first}:T{last}",
        FormulaRule(formula=[f'LEFT($T{first},4)="Over"'], font=S.f(10, True, S.RED_ALERT)))
    ws.conditional_formatting.add(
        f"N{first}:N{last}",
        CellIsRule(operator="greaterThan", formula=["0"], font=S.f(10, True, S.AMBER)))
    S.print_setup(ws, title_rows=f"{head}:{head}")
    return dict(head=head, first=first, last=last, total_row=total_row,
                alloc_row=alloc_row)


def _rng(sheet: str, col: str, first: int, last: int) -> str:
    return f"'{sheet}'!${col}${first}:${col}${last}"


# ── 2. Group asks ───────────────────────────────────────────────────────────
def build_asks(wb: Workbook, bank: R.Bank) -> dict:
    ws = wb.create_sheet(SH_ASK)
    S.sheet_title(
        ws, "Group asks",
        f"One band of {INTAKE_ROWS} rows per group. Copy A6:R205 from a group's "
        f"{T.SH_ASK} sheet and paste values into column B of that group's band. "
        "The computed columns on the right are recomputed centrally — leave them.")
    heads = ["Group"] + PASTE + COMPUTED
    widths = ([26] + [6, 22, 24, 24, 22, 28, 16, 16, 13, 13, 10, 7, 7, 7, 7, 13, 26, 34]
              + [13, 7, 12, 10, 15, 15, 12, 34])
    S.header_row(ws, HEAD_ROW, heads, widths)

    n = len(bank.groups)
    for i, g in enumerate(bank.groups):
        bf = band_first(i)
        for j in range(INTAKE_ROWS):
            r = bf + j
            gc = ws.cell(row=r, column=1,
                         value=f'=IF(${K["Job title"]}{r}="","","{g}")')
            gc.font = S.f(9, True, S.NAVY) if j == 0 else S.SMALL
            for col in range(2, len(PASTE) + 2):        # the paste target
                cell = S.input_cell(ws.cell(row=r, column=col))
                if K["Q1"] <= get_column_letter(col) <= K["New Asks"]:
                    cell.number_format = S.FTE

            level = f'${K["Career Level"]}{r}'
            ws.cell(row=r, column=_col(K["Category"]), value=(
                f'=IF(${K["Driver"]}{r}="","",IFERROR(INDEX(DriverCategory,'
                f'MATCH(${K["Driver"]}{r},DriverList,0)),"{R.RANK_CATEGORIES[-1][0]}"))'))
            ws.cell(row=r, column=_col(K["Rank"]), value=(
                f'=IF(${K["Driver"]}{r}="","",IFERROR(INDEX(DriverRank,'
                f'MATCH(${K["Driver"]}{r},DriverList,0)),99))'))
            ws.cell(row=r, column=_col(K["Rate"]), value=(
                f'=IF({level}="",0,{rate_of(level)})')).number_format = S.MONEY
            ws.cell(row=r, column=_col(K["One-off"]), value=(
                f'=IFERROR(INDEX(LevelOneOff,MATCH({level},LevelList,0)),0)')
            ).number_format = S.MONEY
            ws.cell(row=r, column=_col(K["Requested run-rate"]), value=(
                f'=IF(${K["Job title"]}{r}="","",'
                f'${K["Rate"]}{r}*N(${K["New Asks"]}{r}))')).number_format = S.MONEY
            # Part-year, quarter by quarter: a Q1 start is paid four quarters.
            phased = "+".join(
                f'N(${K[q]}{r})*{(4 - x) / 4}' for x, q in enumerate(R.QUARTERS))
            ws.cell(row=r, column=_col(K[f"Requested {R.PLAN_YEAR} cost"]), value=(
                f'=IF(${K["Job title"]}{r}="","",'
                f'${K["Rate"]}{r}*({phased})'
                f'+${K["One-off"]}{r}*N(${K["New Asks"]}{r}))')).number_format = S.MONEY
            ws.cell(row=r, column=_col(K["Sort key"]), value=(
                f'=IF(${K["Job title"]}{r}="","",${K["Rank"]}{r}*1000000+ROW())')
            ).number_format = S.COUNT

            problem = "&".join(
                f'IF({test},", {word}","")' for test, word in [
                    (f'${K["Division"]}{r}=""', "no division"),
                    (f'${K["Unit"]}{r}=""', "no unit"),
                    (f'{level}=""', "no career level"),
                    (f'N(${K["New Asks"]}{r})<=0', "no quarter split"),
                    (f'${K["Driver"]}{r}=""', "no driver"),
                    (f'AND({level}<>"",COUNTIF(CareerLevels,{level})=0)',
                     "a career level that is not on the ladder"),
                    (f'AND(${K["Driver"]}{r}<>"",'
                     f'COUNTIF(DriverList,${K["Driver"]}{r})=0)',
                     "a driver that is not on the list"),
                    (f'ABS(N(${K["New Asks"]}{r})-({"+".join(f"N(${K[q]}{r})" for q in R.QUARTERS)}))>0.001',
                     "a total that does not match its quarters"),
                ])
            ws.cell(row=r, column=_col(K["Problem"]), value=(
                f'=IF(${K["Job title"]}{r}="","",MID({problem},3,300))')
            ).font = S.f(9, color=S.RED_ALERT)
            for col in range(len(PASTE) + 2, len(PASTE) + len(COMPUTED) + 2):
                if get_column_letter(col) != K["Problem"]:
                    ws.cell(row=r, column=col).font = S.BODY
        S.band(ws, bf, len(heads))

    last = last_row(n)
    ws.conditional_formatting.add(
        f'A{FIRST_ROW}:{K["Problem"]}{last}',
        FormulaRule(formula=[f'${K["Problem"]}{FIRST_ROW}<>""'],
                    fill=S.PatternFill("solid", fgColor="FDE7E9")))
    ws.auto_filter.ref = f'A{HEAD_ROW}:{K["Problem"]}{last}'

    for label, col in [
            ("BookGroup", "Group"), ("BookDivision", "Division"),
            ("BookUnit", "Unit"), ("BookTitle", "Job title"),
            ("BookLevel", "Career Level"), ("BookFamily", "Job Family"),
            ("BookWorkerType", "Worker type"), ("BookTotal", "New Asks"),
            ("BookQ1", "Q1"), ("BookQ2", "Q2"), ("BookQ3", "Q3"), ("BookQ4", "Q4"),
            ("BookCategory", "Category"), ("BookRank", "Rank"), ("BookRate", "Rate"),
            ("BookOneOff", "One-off"), ("BookRunRate", "Requested run-rate"),
            ("BookInYear", f"Requested {R.PLAN_YEAR} cost"),
            ("BookKey", "Sort key"), ("BookProblem", "Problem"),
            ("BookDriver", "Driver"), ("BookRef", "Ref")]:
        _name(wb, SH_ASK, label, f"${K[col]}${FIRST_ROW}:${K[col]}${last}")

    S.print_setup(ws, title_rows=f"{HEAD_ROW}:{HEAD_ROW}")
    return dict(first=FIRST_ROW, last=last)


def _col(letter: str) -> int:
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(letter)


# ── 3. Group returns ────────────────────────────────────────────────────────
def build_returns(wb: Workbook, bank: R.Bank) -> dict:
    ws = wb.create_sheet(SH_RET)
    S.sheet_title(
        ws, "Group returns",
        f"The block at the foot of each group's {T.SH_POS} sheet, one band of "
        f"{RETURN_ROWS} rows per group. This is where the bridge gets its exits "
        "and where the mandated-seat share comes from.")
    S.header_row(ws, HEAD_ROW, ["Group"] + RETURN_COLS,
                 [26, 30, 15, 10, 12, 14, 14, 15])
    n = len(bank.groups)
    for i, g in enumerate(bank.groups):
        rf = ret_first(i)
        for j in range(RETURN_ROWS):
            r = rf + j
            # Keyed off the numbers, not off the division name: a group office
            # row can legitimately arrive without one, and dropping it would
            # quietly lose seats from the bridge.
            gc = ws.cell(row=r, column=1, value=(
                f'=IF(COUNT(${RET["Current Capacity"]}{r}:'
                f'${RET["Run-rate cost"]}{r})=0,"","{g}")'))
            gc.font = S.f(9, True, S.NAVY) if j == 0 else S.SMALL
            for col in range(2, len(RETURN_COLS) + 2):
                cell = S.input_cell(ws.cell(row=r, column=col))
                cell.number_format = S.MONEY if col == 8 else S.FTE
            ws.cell(row=r, column=2).number_format = "General"
        S.band(ws, rf, len(RETURN_COLS) + 1)
    last = ret_last(n)
    for label, col in [("RetGroup", "Group"), ("RetDivision", "Division"),
                       ("RetCurrent", "Current Capacity"), ("RetExits", "Exits"),
                       ("RetAsks", "New Asks"), ("RetNew", "New Capacity"),
                       ("RetMandated", "Mandated seats"),
                       ("RetCost", "Run-rate cost")]:
        _name(wb, SH_RET, label, f"${RET[col]}${FIRST_ROW}:${RET[col]}${last}")
    S.note_line(ws, last + 2,
                "Exits are seats the group has marked Exit on its own capacity "
                "sheet. They come out of the establishment before anything is "
                "added back, which is what makes the bridge balance.",
                cols=len(RETURN_COLS) + 1)
    S.print_setup(ws, title_rows=f"{HEAD_ROW}:{HEAD_ROW}")
    return dict(first=FIRST_ROW, last=last)


# ── 4. Challenge ────────────────────────────────────────────────────────────
def build_challenge(wb: Workbook, bank: R.Bank) -> dict:
    ws = wb.create_sheet(SH_CHAL)
    S.sheet_title(
        ws, "Challenge",
        "Requested against approved, line by line. Set a decision on every line; "
        "only 'Approve fewer' needs a number. Approved is what rolls forward.")
    heads = ["Group", "Ref", "Division", "Unit", "Job title", "Career Level",
             "Category", "Rank", "Requested", "Requested start",
             f"Requested {R.PLAN_YEAR} cost",
             "Decision", "Approved", "Approved start", "Reason",
             "Effective", "Effective start", f"Approved {R.PLAN_YEAR} cost",
             "Approved run-rate", "Cost variance", "Approved line"]
    widths = [26, 6, 22, 24, 28, 16, 13, 7, 11, 13, 15,
              20, 11, 13, 38, 11, 13, 15, 14, 13, 12]
    S.header_row(ws, HEAD_ROW, heads, widths)

    n = len(bank.groups)
    first, last = FIRST_ROW, last_row(n)
    for r in range(first, last + 1):
        for col, src in [(1, "Group"), (2, "Ref"), (3, "Division"), (4, "Unit"),
                         (5, "Job title"), (6, "Career Level"), (7, "Category"),
                         (8, "Rank"), (9, "New Asks"),
                         (11, f"Requested {R.PLAN_YEAR} cost")]:
            ws.cell(row=r, column=col, value=(
                f'=IF(INDEX(BookTitle,{r - first + 1})="","",'
                f'INDEX(Book{_book(src)},{r - first + 1}))')).font = S.LINK
        ws.cell(row=r, column=9).number_format = S.FTE
        ws.cell(row=r, column=11).number_format = S.MONEY
        # The earliest quarter the group asked for, which is the start the
        # centre challenges against.
        ws.cell(row=r, column=10, value=(
            f'=IF($I{r}="","",'
            + "".join(f'IF(INDEX(BookQ{q[1]},{r - first + 1})>0,"{q}",'
                      for q in R.QUARTERS)
            + '""' + ")" * len(R.QUARTERS) + ")"))
        for col in range(12, 16):
            S.input_cell(ws.cell(row=r, column=col))
        ws.cell(row=r, column=13).number_format = S.FTE

        # What the decision means, in numbers.
        ws.cell(row=r, column=16, value=(
            f'=IF($E{r}="","",IF($L{r}="Decline",0,'
            f'IF($L{r}="Approve fewer",IF($M{r}="","",$M{r}),'
            f'IF($L{r}="","",$I{r}))))')).number_format = S.FTE
        ws.cell(row=r, column=17, value=f'=IF($P{r}="","",IF($N{r}<>"",$N{r},$J{r}))')
        rate = f"INDEX(BookRate,{r - first + 1})"
        one_off = f"INDEX(BookOneOff,{r - first + 1})"
        ws.cell(row=r, column=18, value=(
            f'=IF(OR($P{r}="",$Q{r}=""),"",'
            f'{rate}*$P{r}*(5-{qidx(f"$Q{r}")})/4+{one_off}*$P{r})')
        ).number_format = S.MONEY
        ws.cell(row=r, column=19, value=(
            f'=IF($P{r}="","",{rate}*$P{r})')).number_format = S.MONEY
        ws.cell(row=r, column=20, value=(
            f'=IF($R{r}="","",$R{r}-$K{r})')).number_format = S.MONEY
        ws.cell(row=r, column=21, value=(
            f'=IF(N($P{r})<=0,"",COUNTIF($P${first}:$P{r},">0"))')
        ).number_format = S.COUNT
        for col in (16, 17, 18, 19, 20, 21):
            ws.cell(row=r, column=col).font = S.BODY

    dv = DataValidation(type="list", formula1="=DecisionList", allow_blank=True,
                        showErrorMessage=True)
    dv.error = "Pick a decision from the list."
    ws.add_data_validation(dv)
    dv.add(f"L{first}:L{last}")
    dvq = DataValidation(type="list", formula1="=QuarterList", allow_blank=True,
                         showErrorMessage=True)
    ws.add_data_validation(dvq)
    dvq.add(f"N{first}:N{last}")

    ws.conditional_formatting.add(
        f"L{first}:L{last}",
        CellIsRule(operator="equal", formula=['"Decline"'],
                   font=S.f(10, True, S.RED_ALERT)))
    ws.conditional_formatting.add(
        f"T{first}:T{last}",
        CellIsRule(operator="lessThan", formula=["0"], font=S.f(10, color=S.OK_GREEN)))
    ws.auto_filter.ref = f"A{HEAD_ROW}:U{last}"

    for label, col in [("ChalGroup", "A"), ("ChalDivision", "C"), ("ChalUnit", "D"),
                       ("ChalTitle", "E"), ("ChalLevel", "F"), ("ChalCategory", "G"),
                       ("ChalRequested", "I"), ("ChalEffective", "P"),
                       ("ChalStart", "Q"), ("ChalInYear", "R"),
                       ("ChalRunRate", "S"), ("ChalSeq", "U")]:
        _name(wb, SH_CHAL, label, f"${col}${first}:${col}${last}")

    S.print_setup(ws, title_rows=f"{HEAD_ROW}:{HEAD_ROW}")
    return dict(first=first, last=last)


def _book(name: str) -> str:
    """The defined name on the asks sheet that carries this column."""
    return {"Group": "Group", "Ref": "Ref", "Division": "Division", "Unit": "Unit",
            "Job title": "Title", "Career Level": "Level", "Category": "Category",
            "Rank": "Rank", "New Asks": "Total",
            f"Requested {R.PLAN_YEAR} cost": "InYear"}[name]


# ── 5. Scenarios, and the engine behind them ────────────────────────────────
LEVERS = [
    ("Demand multiplier", "Scales every requested position. 1.10 asks what a "
     "tenth more demand would cost.", S.PCT),
    ("Share of envelope released", "How much of the bank envelope this case "
     "funds. The ranked cut falls where this runs out.", S.PCT),
    ("Timing shift (quarters)", "Pushes every start back by this many quarters. "
     "It changes the cash, not the headcount — see the note below.", None),
    ("Attrition rate", "Used for expected departures, not for the cut.", S.PCT),
    ("Salary inflation", "Applied to the rate card for this case.", S.PCT),
    ("Productivity multiplier", "1.0 takes the groups' own productivity "
     "assumption as given. 1.5 assumes half as much again lands, which shrinks "
     "the ask; 0.5 assumes half of it does not.", None),
]
ENG_BLOCK = 5                   # FTE, run-rate, cumulative, funded, in-year cash


def eng_col(scenario: int, offset: int) -> str:
    return get_column_letter(4 + scenario * ENG_BLOCK + offset)


def build_scenarios(wb: Workbook, bank: R.Bank) -> dict:
    ws = wb.create_sheet(SH_SCEN)
    row = S.sheet_title(
        ws, "Scenarios",
        "Four cases over the same book of asks, computed side by side. Change a "
        "lever and every number below moves with it.")
    ws.column_dimensions["A"].width = 40
    for c in "BCDE":
        ws.column_dimensions[c].width = 16
    ws.column_dimensions["F"].width = 58

    names = [s[0] for s in R.SCENARIOS]
    ws.cell(row=row, column=1, value="The levers").font = S.H1
    row += 1
    S.header_row(ws, row, ["Lever"] + names + ["What it does"],
                 [40, 16, 16, 16, 16, 58])
    ws.freeze_panes = None
    lever_first = row + 1
    for i, (label, why, fmt) in enumerate(LEVERS):
        r = lever_first + i
        ws.cell(row=r, column=1, value=label).font = S.H2
        for k, sc in enumerate(R.SCENARIOS):
            value = [sc[1], sc[2], sc[3], sc[4], sc[5], sc[6]][i]
            S.input_cell(ws.cell(row=r, column=2 + k, value=value), fmt)
        ws.cell(row=r, column=6, value=why).font = S.SMALL
        ws.cell(row=r, column=6).alignment = S.WRAP
        ws.row_dimensions[r].height = 26
    lever_rows = {name: lever_first + i for i, (name, _, _) in enumerate(LEVERS)}
    row = lever_first + len(LEVERS) + 1

    ws.cell(row=row, column=1, value="What each case funds").font = S.H1
    row += 1
    S.header_row(ws, row, ["Measure"] + names + ["Read it this way"],
                 [40, 16, 16, 16, 16, 58])
    ws.freeze_panes = None
    return dict(ws=ws, lever_first=lever_first, lever_rows=lever_rows,
                metrics_first=row + 1, names=names,
                first=FIRST_ROW, last=last_row(len(bank.groups)))


def build_engine(wb: Workbook, bank: R.Bank, sc: dict) -> dict:
    """One row per submitted line, computed under all four cases in parallel.
    Nothing to read here; it is the arithmetic the Scenarios sheet summarises."""
    ws = wb.create_sheet(SH_ENG)
    S.sheet_title(ws, "Engine",
                  "Working sheet. Every submitted line under all four cases. "
                  "Nothing here is typed in, and nothing here needs reading.")
    first, last = sc["first"], sc["last"]
    scn = f"'{SH_SCEN}'!"
    lr = sc["lever_rows"]

    heads = ["Sort key", "Group", "Priority"]
    for name in sc["names"]:
        heads += [f"{name} positions", f"{name} run-rate", f"{name} cumulative",
                  f"{name} funded", f"{name} cash"]
    S.header_row(ws, HEAD_ROW, heads,
                 [12, 26, 14] + [13] * (ENG_BLOCK * len(sc["names"])))

    for r in range(first, last + 1):
        pos = r - first + 1
        ws.cell(row=r, column=1, value=f'=IF(INDEX(BookTitle,{pos})="","",'
                                      f"INDEX(BookKey,{pos}))")
        ws.cell(row=r, column=2, value=f'=IF(INDEX(BookTitle,{pos})="","",'
                                      f"INDEX(BookGroup,{pos}))")
        ws.cell(row=r, column=3, value=f'=IF(INDEX(BookTitle,{pos})="","",'
                                      f"INDEX(BookCategory,{pos}))")
        for k in range(len(sc["names"])):
            L = get_column_letter(2 + k)
            c = [eng_col(k, i) for i in range(ENG_BLOCK)]
            base = f"{scn}${L}$"
            scale = (f"{base}{lr['Demand multiplier']}"
                     f"/(1+ProductivityDefault*({base}{lr['Productivity multiplier']}-1))")
            rate = (f"INDEX(BookRate,{pos})*(1+{base}{lr['Salary inflation']})")
            shift = f"{base}{lr['Timing shift (quarters)']}"
            phased = "+".join(
                f"INDEX(BookQ{q[1]},{pos})*(5-MIN(4,{i + 1}+{shift}))/4"
                for i, q in enumerate(R.QUARTERS))
            ws.cell(row=r, column=4 + k * ENG_BLOCK, value=(
                f'=IF(INDEX(BookTitle,{pos})="","",INDEX(BookTotal,{pos})*{scale})')
            ).number_format = S.FTE
            ws.cell(row=r, column=5 + k * ENG_BLOCK, value=(
                f'=IF(${c[0]}{r}="","",{rate}*${c[0]}{r})')).number_format = S.MONEY
            ws.cell(row=r, column=6 + k * ENG_BLOCK, value=(
                f'=IF(${c[1]}{r}="","",SUMIFS(${c[1]}${first}:${c[1]}${last},'
                f'$A${first}:$A${last},"<="&$A{r}))')).number_format = S.MONEY
            ws.cell(row=r, column=7 + k * ENG_BLOCK, value=(
                f'=IF(${c[1]}{r}="","",IF(${c[2]}{r}<=BankCostEnvelope*'
                f"{base}{lr['Share of envelope released']},1,0))")
            ).number_format = S.COUNT
            ws.cell(row=r, column=8 + k * ENG_BLOCK, value=(
                f'=IF(${c[0]}{r}="","",{rate}*({phased})*{scale}'
                f"+INDEX(BookOneOff,{pos})*${c[0]}{r})")).number_format = S.MONEY
    ws.sheet_properties.tabColor = S.GREY_BG
    S.print_setup(ws)
    return dict(first=first, last=last)


def finish_scenarios(wb: Workbook, bank: R.Bank, sc: dict, env: dict) -> dict:
    """Written after the engine exists, because every measure reads from it."""
    ws = sc["ws"]
    first, last, names = sc["first"], sc["last"], sc["names"]
    row = sc["metrics_first"]
    eng = f"'{SH_ENG}'!"
    envt = f"'{SH_ENV}'!"
    opening = f"{envt}$B${env['total_row']}"
    lr = sc["lever_rows"]

    def col(k, off):
        return f"{eng}${eng_col(k, off)}${first}:${eng_col(k, off)}${last}"

    measures = []
    for k in range(len(names)):
        fte, runrate, funded, cash = (col(k, 0), col(k, 1), col(k, 3), col(k, 4))
        L = get_column_letter(2 + k)
        measures.append({
            "Positions requested (after the demand lever)": (f"=SUM({fte})", S.FTE),
            "Lines funded": (f"=COUNTIFS({funded},1)", S.COUNT),
            "Positions funded": (f"=SUMIFS({fte},{funded},1)", S.FTE),
            "Run-rate cost funded (SAR '000)":
                (f"=SUMIFS({runrate},{funded},1)", S.MONEY),
            f"{R.PLAN_YEAR} cash cost of those (SAR '000)":
                (f"=SUMIFS({cash},{funded},1)", S.MONEY),
            "Envelope released (SAR '000)":
                (f"=BankCostEnvelope*${L}${lr['Share of envelope released']}", S.MONEY),
            "Positions not funded":
                (f"=SUM({fte})-SUMIFS({fte},{funded},1)", S.FTE),
            "Regulatory positions funded":
                (f'=SUMIFS({fte},{funded},1,{eng}$C${first}:$C${last},"Regulatory")',
                 S.FTE),
            "Closing establishment":
                (f"={opening}-SUM(RetExits)+SUMIFS({fte},{funded},1)", S.FTE),
            "Mandated share of it":
                (f"=IFERROR(SUM(RetMandated)/({opening}-SUM(RetExits)"
                 f"+SUMIFS({fte},{funded},1)),0)", S.PCT),
            "Against the headcount envelope":
                (f'=IF(SUMIFS({fte},{funded},1)-SUM(RetExits)<=BankHeadEnvelope,'
                 f'"Within","Over by "&TEXT(SUMIFS({fte},{funded},1)-SUM(RetExits)'
                 f'-BankHeadEnvelope,"#,##0.0"))', None),
            "Expected departures at this rate":
                (f"=ROUND({envt}$C${env['total_row']}*${L}${lr['Attrition rate']},0)",
                 S.COUNT),
        })
    notes = {
        "Positions requested (after the demand lever)":
            "What the groups asked for, moved by this case's demand assumption.",
        "Lines funded": "How many of the requested lines survive the ranked cut.",
        "Positions funded": "The number that matters: what the bank would hire.",
        "Run-rate cost funded (SAR '000)":
            "The full-year cost the bank carries. This is what the cut is made against.",
        f"{R.PLAN_YEAR} cash cost of those (SAR '000)":
            "What the same positions cost in the plan year, at this case's start quarters.",
        "Envelope released (SAR '000)": "The ceiling this case is cut against.",
        "Positions not funded": "What falls out. Read it with the priority table below.",
        "Regulatory positions funded":
            "Regulatory work ranks first, so this should hold in every case.",
        "Closing establishment":
            "Opening establishment, less the seats the groups are giving up, plus "
            "what this case funds.",
        "Mandated share of it":
            "Seats the groups have marked as having to be Saudi, over the closing "
            "establishment. New positions carry no mandate yet, so this drifts down "
            "as the bank grows — that is the point of watching it.",
        "Against the headcount envelope": "Funded positions less exits.",
        "Expected departures at this rate":
            "Not part of the cut — a backfill is already a line in the book.",
    }
    for label in measures[0]:
        ws.cell(row=row, column=1, value=label).font = S.H2
        for k in range(len(names)):
            formula, fmt = measures[k][label]
            c = ws.cell(row=row, column=2 + k, value=formula)
            if fmt:
                c.number_format = fmt
            c.font = S.BODY
        ws.cell(row=row, column=6, value=notes[label]).font = S.SMALL
        ws.cell(row=row, column=6).alignment = S.WRAP
        for c in range(1, 7):
            ws.cell(row=row, column=c).border = S.BOX
        row += 1
    metrics_last = row - 1
    runrate_row = sc["metrics_first"] + 3
    cash_row = sc["metrics_first"] + 4
    fte_row = sc["metrics_first"] + 2
    row += 1

    S.note_line(
        ws, row,
        "The cut is made against run-rate cost, so pushing starts back a quarter "
        f"does not buy a single extra position — it buys cash in {R.PLAN_YEAR} and "
        "leaves the commitment where it was. Read the run-rate row and the cash row "
        "together; they answer different questions and only one of them is a saving.",
        cols=6)
    row += 2

    ws.cell(row=row, column=1, value="Funded positions by priority").font = S.H1
    row += 1
    S.header_row(ws, row, ["Priority"] + names + [""], [40, 16, 16, 16, 16, 58])
    ws.freeze_panes = None
    cat_first = row + 1
    for i, (name, rank, why) in enumerate(R.RANK_CATEGORIES):
        r = cat_first + i
        ws.cell(row=r, column=1, value=f"{rank}. {name}").font = S.BODY
        for k in range(len(names)):
            ws.cell(row=r, column=2 + k, value=(
                f'=SUMIFS({col(k, 0)},{col(k, 3)},1,'
                f'{eng}$C${first}:$C${last},"{name}")')).number_format = S.FTE
        ws.cell(row=r, column=6, value=why).font = S.SMALL
        ws.cell(row=r, column=6).alignment = S.WRAP
        for c in range(1, 7):
            ws.cell(row=r, column=c).border = S.BOX
    cat_last = cat_first + len(R.RANK_CATEGORIES) - 1
    row = cat_last + 2

    ws.cell(row=row, column=1, value="Funded run-rate cost by group (SAR '000)").font = S.H1
    row += 1
    S.header_row(ws, row, ["Group"] + names + [""], [40, 16, 16, 16, 16, 58])
    ws.freeze_panes = None
    grp_first = row + 1
    for i, g in enumerate(bank.groups):
        r = grp_first + i
        q = g.replace('"', '""')
        ws.cell(row=r, column=1, value=g).font = S.BODY
        for k in range(len(names)):
            ws.cell(row=r, column=2 + k, value=(
                f'=SUMIFS({col(k, 1)},{col(k, 3)},1,'
                f'{eng}$B${first}:$B${last},"{q}")')).number_format = S.MONEY
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = S.BOX
    grp_last = grp_first + len(bank.groups) - 1

    chart = BarChart()
    chart.type = "col"
    chart.title = "Run-rate cost funded by case (SAR '000)"
    chart.legend = None
    chart.height, chart.width = 8, 16
    chart.add_data(Reference(ws, min_col=2, max_col=5, min_row=runrate_row,
                             max_row=runrate_row), from_rows=True,
                   titles_from_data=False)
    chart.set_categories(Reference(ws, min_col=2, max_col=5,
                                   min_row=sc["metrics_first"] - 1,
                                   max_row=sc["metrics_first"] - 1))
    chart.series[0].graphicalProperties.solidFill = S.NAVY
    ws.add_chart(chart, f"H{sc['lever_first']}")

    stack = BarChart()
    stack.type = "col"
    stack.grouping = "clustered"
    stack.title = "Funded positions by priority"
    stack.height, stack.width = 8, 16
    stack.add_data(Reference(ws, min_col=1, max_col=5, min_row=cat_first - 1,
                             max_row=cat_last), titles_from_data=True)
    stack.set_categories(Reference(ws, min_col=1, min_row=cat_first, max_row=cat_last))
    ws.add_chart(stack, f"H{cat_first - 2}")

    S.print_setup(ws)
    return dict(names=names, metrics_first=sc["metrics_first"],
                metrics_last=metrics_last, runrate_row=runrate_row,
                cash_row=cash_row, fte_row=fte_row,
                cat_first=cat_first, cat_last=cat_last,
                grp_first=grp_first, grp_last=grp_last)


# ── 6. Dashboard ────────────────────────────────────────────────────────────
def build_dashboard(wb: Workbook, bank: R.Bank, env: dict) -> dict:
    ws = wb.create_sheet(SH_DASH)
    row = S.sheet_title(
        ws, "Dashboard",
        f"The {R.BASE_YEAR} to {R.PLAN_YEAR} bridge, the groups against their "
        "envelopes, demand by driver, and how the cost phases across the year.")
    for col, w in zip("ABCDEFGH", [40, 15, 15, 15, 15, 15, 15, 15]):
        ws.column_dimensions[col].width = w

    charts = {"row": 5}

    def place(chart, height_cm: float, width_cm: float = 18) -> None:
        chart.height, chart.width = height_cm, width_cm
        ws.add_chart(chart, f"J{charts['row']}")
        charts["row"] += int(height_cm * 1.9) + 3

    envt = f"'{SH_ENV}'!"
    et = env["total_row"]

    # ── The bridge ─────────────────────────────────────────────────────────
    ws.cell(row=row, column=1,
            value=f"{R.BASE_YEAR} to {R.PLAN_YEAR} establishment bridge").font = S.H1
    ws.cell(row=row, column=6,
            value="Approved seats, not people. A seat only leaves the "
                  "establishment if a group marked it Exit.").font = S.NOTE
    row += 1
    S.header_row(ws, row, ["Step", "Change", "Running total", "Base", "Bar"],
                 [40, 15, 15, 15, 15])
    ws.freeze_panes = None
    b_first = row + 1
    steps = [
        (f"Opening {R.BASE_YEAR} approved establishment", f"={envt}$B${et}"),
        ("less seats the groups are giving up", "=-SUM(RetExits)"),
        ("plus positions approved", "=SUM(ChalEffective)"),
        (f"Closing {R.PLAN_YEAR} approved establishment", None),
    ]
    for i, (label, formula) in enumerate(steps):
        r = b_first + i
        ws.cell(row=r, column=1, value=label).font = S.H2 if i in (0, 3) else S.BODY
        if formula:
            ws.cell(row=r, column=2, value=formula).number_format = S.FTE
        if i == 0:
            ws.cell(row=r, column=3, value=f"=$B{r}")
            ws.cell(row=r, column=4, value=0)
            ws.cell(row=r, column=5, value=f"=$B{r}")
        elif i < 3:
            ws.cell(row=r, column=3, value=f"=$C{r - 1}+$B{r}")
            ws.cell(row=r, column=4, value=f"=MIN($C{r - 1},$C{r})")
            ws.cell(row=r, column=5, value=f"=ABS($B{r})")
        else:
            ws.cell(row=r, column=2, value=f"=$C{r - 1}")
            ws.cell(row=r, column=3, value=f"=$C{r - 1}")
            ws.cell(row=r, column=4, value=0)
            ws.cell(row=r, column=5, value=f"=$C{r - 1}")
        for col in range(2, 6):
            ws.cell(row=r, column=col).number_format = S.FTE
            ws.cell(row=r, column=col).border = S.BOX
        ws.cell(row=r, column=1).border = S.BOX
    b_last = b_first + len(steps) - 1
    S.band(ws, b_first, 5)
    S.band(ws, b_last, 5)
    row = b_last + 1
    for label, formula in [
        ("Net establishment change against the bank envelope",
         f'=IF($C${b_last}-$B${b_first}<=BankHeadEnvelope,"Within the envelope",'
         f'"Over the envelope by "'
         f'&TEXT($C${b_last}-$B${b_first}-BankHeadEnvelope,"#,##0.0"))'),
        ("Mandated seats as a share of the closing establishment",
         f"=IFERROR(SUM(RetMandated)/$C${b_last},0)"),
    ]:
        ws.cell(row=row, column=1, value=label).font = S.BODY_DIM
        c = ws.cell(row=row, column=2, value=formula)
        c.font = S.BODY
        if "share" in label:
            c.number_format = S.PCT
        row += 1
    row += 1

    wf = BarChart()
    wf.type = "col"
    wf.grouping = "stacked"
    wf.overlap = 100
    wf.title = f"{R.BASE_YEAR} to {R.PLAN_YEAR} establishment"
    wf.legend = None
    wf.add_data(Reference(ws, min_col=4, min_row=b_first, max_row=b_last),
                titles_from_data=False)
    wf.add_data(Reference(ws, min_col=5, min_row=b_first, max_row=b_last),
                titles_from_data=False)
    wf.set_categories(Reference(ws, min_col=1, min_row=b_first, max_row=b_last))
    wf.series[0].graphicalProperties.noFill = True
    wf.series[1].graphicalProperties.solidFill = S.NAVY
    place(wf, 8)

    # ── League table ───────────────────────────────────────────────────────
    ws.cell(row=row, column=1, value="Groups against their envelopes").font = S.H1
    row += 1
    S.header_row(ws, row, ["Group", "Requested", "Approved", "Approved run-rate",
                           "Cost envelope", "Share of envelope", "Status"],
                 [40, 13, 13, 16, 15, 15, 22])
    ws.freeze_panes = None
    lg_first = row + 1
    for i, g in enumerate(bank.groups):
        r = lg_first + i
        er = env["first"] + i
        ws.cell(row=r, column=1, value=f"={envt}$A${er}").font = S.BODY
        ws.cell(row=r, column=2, value=f"={envt}$O${er}").number_format = S.FTE
        ws.cell(row=r, column=3, value=f"={envt}$R${er}").number_format = S.FTE
        ws.cell(row=r, column=4, value=f"={envt}$S${er}").number_format = S.MONEY
        ws.cell(row=r, column=5, value=f"={envt}$F${er}").number_format = S.MONEY
        ws.cell(row=r, column=6, value=f"=IFERROR($D{r}/$E{r},0)").number_format = S.PCT
        ws.cell(row=r, column=7, value=f"={envt}$T${er}")
        for col in range(1, 8):
            ws.cell(row=r, column=col).border = S.BOX
    lg_last = lg_first + len(bank.groups) - 1
    ws.conditional_formatting.add(
        f"F{lg_first}:F{lg_last}",
        CellIsRule(operator="greaterThan", formula=["1"], font=S.f(10, True, S.RED_ALERT)))
    row = lg_last + 2

    league = BarChart()
    league.type = "bar"
    league.title = f"Approved {R.PLAN_YEAR} run-rate against envelope (SAR '000)"
    league.add_data(Reference(ws, min_col=4, max_col=5, min_row=lg_first - 1,
                              max_row=lg_last), titles_from_data=True)
    league.set_categories(Reference(ws, min_col=1, min_row=lg_first, max_row=lg_last))
    league.series[0].graphicalProperties.solidFill = S.NAVY
    league.series[1].graphicalProperties.solidFill = "C7D3E3"
    league.x_axis.scaling.orientation = "maxMin"
    league.y_axis.crosses = "max"
    place(league, 10)

    # ── Demand by driver ───────────────────────────────────────────────────
    ws.cell(row=row, column=1, value="Demand by driver — asked against approved").font = S.H1
    row += 1
    S.header_row(ws, row, ["Driver", "Priority", "Requested", "Approved",
                           "Not funded"], [40, 15, 13, 13, 13])
    ws.freeze_panes = None
    d_first = row + 1
    n = len(bank.groups)
    ch_fte = _rng(SH_CHAL, "P", FIRST_ROW, last_row(n))
    for i, (name, cat) in enumerate(R.ASK_DRIVERS):
        r = d_first + i
        q = name.replace('"', '""')
        ws.cell(row=r, column=1, value=name).font = S.BODY
        ws.cell(row=r, column=2, value=cat).font = S.SMALL
        ws.cell(row=r, column=3,
                value=f'=SUMIFS(BookTotal,BookDriver,"{q}")').number_format = S.FTE
        ws.cell(row=r, column=4,
                value=f'=SUMIFS({ch_fte},BookDriver,"{q}")').number_format = S.FTE
        ws.cell(row=r, column=5, value=f"=$C{r}-$D{r}").number_format = S.FTE
        for col in range(1, 6):
            ws.cell(row=r, column=col).border = S.BOX
    d_last = d_first + len(R.ASK_DRIVERS) - 1
    row = d_last + 2

    drv = BarChart()
    drv.type = "bar"
    drv.title = "Requested against approved, by driver"
    drv.add_data(Reference(ws, min_col=3, max_col=4, min_row=d_first - 1,
                           max_row=d_last), titles_from_data=True)
    drv.set_categories(Reference(ws, min_col=1, min_row=d_first, max_row=d_last))
    drv.series[0].graphicalProperties.solidFill = "C7D3E3"
    drv.series[1].graphicalProperties.solidFill = S.NAVY
    drv.x_axis.scaling.orientation = "maxMin"
    drv.y_axis.crosses = "max"
    place(drv, 14)

    # ── Cost phasing ───────────────────────────────────────────────────────
    ws.cell(row=row, column=1,
            value=f"How the approved {R.PLAN_YEAR} cash cost phases").font = S.H1
    row += 1
    cats = [c for c, _, _ in R.RANK_CATEGORIES]
    S.header_row(ws, row, ["Start quarter"] + cats + ["Total"],
                 [40] + [15] * len(cats) + [15])
    ws.freeze_panes = None
    p_first = row + 1
    ch_q = _rng(SH_CHAL, "Q", FIRST_ROW, last_row(n))
    ch_cost = _rng(SH_CHAL, "R", FIRST_ROW, last_row(n))
    for i, q in enumerate(R.QUARTERS):
        r = p_first + i
        ws.cell(row=r, column=1, value=q).font = S.BODY
        for j, cat in enumerate(cats):
            ws.cell(row=r, column=2 + j, value=(
                f'=SUMIFS({ch_cost},{ch_q},"{q}",BookCategory,"{cat}")')
            ).number_format = S.MONEY
        letter = get_column_letter(1 + len(cats))
        ws.cell(row=r, column=2 + len(cats),
                value=f"=SUM($B{r}:${letter}{r})").number_format = S.MONEY
        for col in range(1, 3 + len(cats)):
            ws.cell(row=r, column=col).border = S.BOX
    p_last = p_first + len(R.QUARTERS) - 1

    phase = BarChart()
    phase.type = "col"
    phase.grouping = "stacked"
    phase.overlap = 100
    phase.title = f"Approved {R.PLAN_YEAR} cash cost by start quarter (SAR '000)"
    phase.add_data(Reference(ws, min_col=2, max_col=1 + len(cats),
                             min_row=p_first - 1, max_row=p_last),
                   titles_from_data=True)
    phase.set_categories(Reference(ws, min_col=1, min_row=p_first, max_row=p_last))
    for series, colour in zip(phase.series, ["24466B", "3E7CB1", "6FA8C7", "C7D3E3"]):
        series.graphicalProperties.solidFill = colour
    place(phase, 9)

    S.print_setup(ws)
    return dict(bridge_first=b_first, bridge_last=b_last, league_first=lg_first,
                league_last=lg_last, driver_first=d_first, driver_last=d_last,
                phase_first=p_first, phase_last=p_last)


# ── 7. Executive summary ────────────────────────────────────────────────────
def build_exec(wb: Workbook, bank: R.Bank, env: dict, scen: dict, dash: dict) -> None:
    ws = wb.create_sheet(SH_EXEC)
    row = S.sheet_title(
        ws, f"{R.PLAN_YEAR} capacity — executive summary",
        "One page. Every number on it comes from the approved column, not the ask.")
    ws.column_dimensions["A"].width = 52
    for c in "BCDE":
        ws.column_dimensions[c].width = 17
    ws.column_dimensions["F"].width = 46

    envt = f"'{SH_ENV}'!"
    et = env["total_row"]
    ws.cell(row=row, column=1, value="The decision in six numbers").font = S.H1
    row += 1
    lines = [
        ("Positions requested by the groups", f"={envt}$O${et}", S.FTE,
         "Twelve submissions, priced centrally on one rate card."),
        ("Positions approved", f"={envt}$R${et}", S.FTE,
         "What the challenge process left standing."),
        ("Seats the groups are giving up", "=SUM(RetExits)", S.FTE,
         "Marked Exit on their own capacity sheets."),
        (f"Approved run-rate cost (SAR '000)", f"={envt}$S${et}", S.MONEY,
         "The full-year cost the bank carries from these decisions."),
        ("Against the bank cost envelope", f"=BankCostEnvelope-{envt}$S${et}", S.MONEY,
         "Positive is headroom; negative is the gap to close."),
        (f"Closing {R.PLAN_YEAR} establishment",
         f"='{SH_DASH}'!$C${dash['bridge_last']}", S.FTE,
         "Opening establishment, less the exits, plus what was approved."),
    ]
    for label, formula, fmt, why in lines:
        ws.cell(row=row, column=1, value=label).font = S.H2
        c = ws.cell(row=row, column=2, value=formula)
        c.number_format = fmt
        c.font = S.f(14, True, S.NAVY)
        ws.cell(row=row, column=3, value=why).font = S.SMALL
        ws.cell(row=row, column=3).alignment = S.WRAP
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=6)
        ws.row_dimensions[row].height = 24
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="The four cases").font = S.H1
    row += 1
    S.header_row(ws, row, ["Case"] + scen["names"], [52, 17, 17, 17, 17])
    ws.freeze_panes = None
    r = row + 1
    for label, src in [("Positions funded", scen["fte_row"]),
                       ("Run-rate cost funded (SAR '000)", scen["runrate_row"]),
                       (f"{R.PLAN_YEAR} cash cost (SAR '000)", scen["cash_row"])]:
        ws.cell(row=r, column=1, value=label).font = S.H2
        for k in range(4):
            letter = get_column_letter(2 + k)
            ws.cell(row=r, column=2 + k,
                    value=f"='{SH_SCEN}'!${letter}${src}").number_format = (
                S.FTE if "Positions" in label else S.MONEY)
        for col in range(1, 6):
            ws.cell(row=r, column=col).border = S.BOX
        r += 1
    row = r + 1

    ws.cell(row=row, column=1, value="Where the money goes").font = S.H1
    row += 1
    S.header_row(ws, row, ["Priority", "Requested", "Approved", "Approved run-rate",
                           "Share of approved cost"], [52, 17, 17, 17, 17])
    ws.freeze_panes = None
    c_first = row + 1
    n = len(bank.groups)
    ch_fte = _rng(SH_CHAL, "P", FIRST_ROW, last_row(n))
    ch_rr = _rng(SH_CHAL, "S", FIRST_ROW, last_row(n))
    for i, (name, rank, _) in enumerate(R.RANK_CATEGORIES):
        r = c_first + i
        ws.cell(row=r, column=1, value=f"{rank}. {name}").font = S.BODY
        ws.cell(row=r, column=2,
                value=f'=SUMIFS(BookTotal,BookCategory,"{name}")').number_format = S.FTE
        ws.cell(row=r, column=3,
                value=f'=SUMIFS({ch_fte},BookCategory,"{name}")').number_format = S.FTE
        ws.cell(row=r, column=4,
                value=f'=SUMIFS({ch_rr},BookCategory,"{name}")').number_format = S.MONEY
        ws.cell(row=r, column=5,
                value=f"=IFERROR($D{r}/SUM($D${c_first}:$D${c_first + 3}),0)"
                ).number_format = S.PCT
        for col in range(1, 6):
            ws.cell(row=r, column=col).border = S.BOX
    c_last = c_first + len(R.RANK_CATEGORIES) - 1
    row = c_last + 2

    for text in [
        "Every figure on this page recalculates from the Challenge sheet. If a "
        "decision changes in the meeting, this page is right again the moment the "
        "cell is typed — there is nothing to refresh and nothing to rebuild.",
        "The rate card behind the cost is illustrative until Finance replaces it. It "
        "sits on the hidden Ref sheet, one row per career level, with basic, housing, "
        "transport, bonus and employer GOSI shown separately so a rate can be "
        "challenged without rebuilding the model.",
    ]:
        S.note_line(ws, row, text, cols=6)
        row += 1
    S.print_setup(ws, landscape=False)


# ── 8. Approved establishment ───────────────────────────────────────────────
def build_establishment(wb: Workbook, bank: R.Bank) -> None:
    ws = wb.create_sheet(SH_EST)
    S.sheet_title(
        ws, f"Approved {R.PLAN_YEAR} establishment",
        "Every approved line, as a flat list for Finance and recruitment. It fills "
        "itself from the Challenge sheet as decisions are made.")
    heads = ["#", "Row", "Group", "Division", "Unit", "Job title", "Career Level",
             "Job Family", "Worker type", "Approved", "Start quarter",
             f"{R.PLAN_YEAR} cash cost", "Run-rate", "Driver"]
    S.header_row(ws, HEAD_ROW, heads,
                 [6, 8, 26, 24, 26, 30, 16, 16, 14, 11, 13, 15, 14, 30])
    ws.column_dimensions["B"].hidden = True

    for i in range(EST_ROWS):
        r = FIRST_ROW + i
        ws.cell(row=r, column=1, value=i + 1).font = S.SMALL
        # One MATCH per row, reused by every column beside it.
        ws.cell(row=r, column=2, value=f'=IFERROR(MATCH($A{r},ChalSeq,0),"")')
        pos = f"$B{r}"
        for col, src in [(3, "ChalGroup"), (4, "ChalDivision"), (5, "ChalUnit"),
                         (6, "ChalTitle"), (7, "ChalLevel"), (8, "BookFamily"),
                         (9, "BookWorkerType"), (10, "ChalEffective"),
                         (11, "ChalStart"), (12, "ChalInYear"), (13, "ChalRunRate"),
                         (14, "BookDriver")]:
            ws.cell(row=r, column=col,
                    value=f'=IF($B{r}="","",INDEX({src},{pos}))')
        for col, fmt in [(10, S.FTE), (12, S.MONEY), (13, S.MONEY)]:
            ws.cell(row=r, column=col).number_format = fmt
    est_last = FIRST_ROW + EST_ROWS - 1
    r = est_last + 1
    ws.cell(row=r, column=6, value="Total").font = S.f(10, True)
    for col in (10, 12, 13):
        letter = get_column_letter(col)
        c = ws.cell(row=r, column=col,
                    value=f"=SUM({letter}{FIRST_ROW}:{letter}{est_last})")
        c.font = S.f(10, True)
        c.number_format = S.FTE if col == 10 else S.MONEY
    ws.auto_filter.ref = f"A{HEAD_ROW}:N{est_last}"
    S.print_setup(ws, title_rows=f"{HEAD_ROW}:{HEAD_ROW}")


# ── 9. Group pages ──────────────────────────────────────────────────────────
def build_group_pages(wb: Workbook, bank: R.Bank, env: dict) -> None:
    from openpyxl.worksheet.pagebreak import Break

    ws = wb.create_sheet(SH_PAGES)
    S.sheet_title(
        ws, "Group pages",
        "One page per group, for sending the answer back. Print the sheet and "
        "every group gets its own sheet of paper.")
    for col, w in zip("ABCDEF", [44, 16, 16, 16, 16, 40]):
        ws.column_dimensions[col].width = w

    n = len(bank.groups)
    envt = f"'{SH_ENV}'!"
    ch_g = _rng(SH_CHAL, "A", FIRST_ROW, last_row(n))
    ch_fte = _rng(SH_CHAL, "P", FIRST_ROW, last_row(n))
    ch_q = _rng(SH_CHAL, "Q", FIRST_ROW, last_row(n))
    ch_cost = _rng(SH_CHAL, "R", FIRST_ROW, last_row(n))
    ch_rr = _rng(SH_CHAL, "S", FIRST_ROW, last_row(n))

    row = 4
    for i, g in enumerate(bank.groups):
        q = g.replace('"', '""')
        er = env["first"] + i
        if i:
            ws.row_breaks.append(Break(id=row - 1))
        ws.cell(row=row, column=1, value=g).font = S.f(16, True, S.NAVY)
        ws.cell(row=row, column=5,
                value=f"Approved {R.PLAN_YEAR} capacity").font = S.SUBTITLE
        S.band(ws, row, 6)
        row += 2
        for label, formula, fmt in [
            ("Positions requested", f"={envt}$O${er}", S.FTE),
            ("Positions approved", f"={envt}$R${er}", S.FTE),
            ("Lines declined", f'=COUNTIFS({ch_g},"{q}",{ch_fte},0)', S.COUNT),
            ("Seats given up", f"={envt}$Q${er}", S.FTE),
            ("Approved run-rate cost (SAR '000)", f"={envt}$S${er}", S.MONEY),
            (f"{R.PLAN_YEAR} cash cost (SAR '000)",
             f'=SUMIFS({ch_cost},{ch_g},"{q}")', S.MONEY),
            ("Cost envelope (SAR '000)", f"={envt}$F${er}", S.MONEY),
            ("Against the envelope", f"={envt}$T${er}", None),
        ]:
            ws.cell(row=row, column=1, value=label).font = S.H2
            c = ws.cell(row=row, column=2, value=formula)
            if fmt:
                c.number_format = fmt
            c.font = S.BODY
            row += 1
        row += 1

        ws.cell(row=row, column=1, value="Approved by priority").font = S.H2
        ws.cell(row=row, column=4, value="Approved by start quarter").font = S.H2
        row += 1
        for label, col in [("Priority", 1), ("Positions", 2), ("Run-rate", 3),
                           ("Quarter", 4), ("Positions", 5), ("Cash cost", 6)]:
            c = ws.cell(row=row, column=col, value=label)
            c.font = S.TH
            c.fill = S.FILL_HEAD
        row += 1
        for j, (name, rank, _) in enumerate(R.RANK_CATEGORIES):
            r = row + j
            ws.cell(row=r, column=1, value=f"{rank}. {name}").font = S.BODY
            ws.cell(row=r, column=2, value=(
                f'=SUMIFS({ch_fte},{ch_g},"{q}",BookCategory,"{name}")')
            ).number_format = S.FTE
            ws.cell(row=r, column=3, value=(
                f'=SUMIFS({ch_rr},{ch_g},"{q}",BookCategory,"{name}")')
            ).number_format = S.MONEY
        for j, quarter in enumerate(R.QUARTERS):
            r = row + j
            ws.cell(row=r, column=4, value=quarter).font = S.BODY
            ws.cell(row=r, column=5, value=(
                f'=SUMIFS({ch_fte},{ch_g},"{q}",{ch_q},"{quarter}")')
            ).number_format = S.FTE
            ws.cell(row=r, column=6, value=(
                f'=SUMIFS({ch_cost},{ch_g},"{q}",{ch_q},"{quarter}")')
            ).number_format = S.MONEY
        row += len(R.RANK_CATEGORIES) + 1
        S.note_line(ws, row,
                    f"Approved is what {g} should now plan and recruit against. "
                    "Anything declined can be brought back at the mid-year review "
                    "with the volumes that have actually landed.", cols=6)
        row += 3
    S.print_setup(ws, landscape=False, fit_width=1, fit_height=0)


# ── Assembly ────────────────────────────────────────────────────────────────
TAB_ORDER = [SH_READ, SH_ENV, SH_ASK, SH_RET, SH_CHAL, SH_SCEN, SH_DASH,
             SH_EXEC, SH_EST, SH_PAGES, SH_ENG, SH_REF]


def build_consolidator(path=None, bank: R.Bank | None = None) -> dict:
    bank = bank or R.Bank()
    path = Path(path or DIST / f"{R.PLAN_YEAR}-Capacity-Consolidator.xlsx")
    wb = Workbook()
    wb.remove(wb.active)

    build_readme(wb, bank)
    env = build_envelopes(wb, bank)
    book = build_asks(wb, bank)
    ret = build_returns(wb, bank)
    chal = build_challenge(wb, bank)
    sc = build_scenarios(wb, bank)
    build_engine(wb, bank, sc)
    scen = finish_scenarios(wb, bank, sc, env)
    dash = build_dashboard(wb, bank, env)
    build_exec(wb, bank, env, scen, dash)
    build_establishment(wb, bank)
    build_group_pages(wb, bank, env)
    build_ref(wb, bank)

    wb._sheets.sort(key=lambda w: TAB_ORDER.index(w.title))
    for ws in wb.worksheets:
        S.protect(ws)
    wb.properties.title = f"{R.PLAN_YEAR} Capacity Exercise - consolidator"
    wb.properties.creator = "Capacity exercise toolkit"
    wb.save(path)
    return dict(path=str(path), env=env, book=book, ret=ret, chal=chal,
                scen=scen, dash=dash, first=FIRST_ROW,
                last=last_row(len(bank.groups)),
                ret_last=ret_last(len(bank.groups)))


if __name__ == "__main__":
    print(build_consolidator()["path"])
