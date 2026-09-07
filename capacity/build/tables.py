"""The lookup tables both workbooks need, written the same way in each.

The template prices an ask and the consolidator prices it again centrally. If
those two ever used different tables the exercise would fall apart on the first
challenge meeting, so the tables are written from one place.
"""
from __future__ import annotations

from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

import refdata as R
import style as S


def _name(wb, sheet: str, name: str, ref: str) -> None:
    wb.defined_names.add(DefinedName(name, attr_text=f"'{sheet}'!{ref}"))


def write_settings(wb, ws, sheet: str, row: int, extras=()) -> int:
    """The rates. Every formula that needs one points at these cells."""
    ws.cell(row=row, column=1, value="Settings").font = S.H1
    row += 1
    settings = [
        ("PlanYear", "Plan year", R.PLAN_YEAR, None),
        ("BaseYear", "Base year", R.BASE_YEAR, None),
        ("AttritionDefault", "Attrition rate (bank default)", R.DEFAULT_ATTRITION, S.PCT),
        ("ProductivityDefault", "Productivity uplift (bank default)",
         R.DEFAULT_PRODUCTIVITY, S.PCT),
        ("SalaryInflation", "Salary inflation", R.DEFAULT_SALARY_INFLATION, S.PCT),
        ("GosiSaudi", "Employer GOSI - Saudi", R.GOSI_SAUDI, S.PCT),
        ("GosiOther", "Employer GOSI - non-Saudi", R.GOSI_NON_SAUDI, S.PCT),
    ] + list(extras)
    for name, label, value, fmt in settings:
        S.label_value(ws, row, label, value, fmt).font = S.INPUT
        _name(wb, sheet, name, f"$B${row}")
        row += 1
    return row + 1


def write_rate_card(wb, ws, sheet: str, row: int) -> int:
    """The rate card, keyed on career level. The one place a cost is stated."""
    ws.cell(row=row, column=1,
            value="Rate card by career level (SAR '000 per annum)").font = S.H1
    ws.cell(row=row, column=7,
            value="Illustrative rates — replace with Finance's own before use.").font = S.NOTE
    row += 1
    S.header_row(ws, row, ["Career level", "Basic", "Housing", "Transport",
                           "Bonus %", "One-off of a hire", "Full-year loaded cost"],
                 [22, 12, 12, 12, 12, 16, 20])
    ws.freeze_panes = None
    first = row + 1
    for i, (level, (b, h, t, bo, one)) in enumerate(R.LEVEL_COST.items()):
        r = first + i
        ws.cell(row=r, column=1, value=level)
        for col, value, fmt in [(2, b, S.MONEY), (3, h, S.MONEY), (4, t, S.MONEY),
                                (5, bo, S.PCT), (6, one, S.MONEY)]:
            ws.cell(row=r, column=col, value=value).number_format = fmt
        # Shown, not hidden: the reader can see what the loaded cost is made of.
        ws.cell(row=r, column=7, value=(
            f"=$B{r}+$C{r}+$D{r}+$B{r}*$E{r}+($B{r}+$C{r})*GosiSaudi")
        ).number_format = S.MONEY
    last = first + len(R.LEVEL_COST) - 1
    for name, col in [("LevelList", "A"), ("LevelBasic", "B"), ("LevelHousing", "C"),
                      ("LevelTransport", "D"), ("LevelBonus", "E"), ("LevelOneOff", "F")]:
        _name(wb, sheet, name, f"${col}${first}:${col}${last}")
    return last + 2


def write_drivers(wb, ws, sheet: str, row: int) -> int:
    ws.cell(row=row, column=1, value="Ask drivers").font = S.H1
    row += 1
    S.header_row(ws, row, ["Driver", "Category", "Rank"], [40, 18, 8])
    ws.freeze_panes = None
    first = row + 1
    for i, (name, cat) in enumerate(R.ASK_DRIVERS):
        r = first + i
        ws.cell(row=r, column=1, value=name)
        ws.cell(row=r, column=2, value=cat)
        ws.cell(row=r, column=3, value=R.rank_of(cat))
    last = first + len(R.ASK_DRIVERS) - 1
    for name, col in [("DriverList", "A"), ("DriverCategory", "B"), ("DriverRank", "C")]:
        _name(wb, sheet, name, f"${col}${first}:${col}${last}")
    return last + 2


def write_lists(wb, ws, sheet: str, row: int, extra=()) -> int:
    """The pick lists, side by side. Returns the next free column."""
    lists = [
        ("QuarterList", R.QUARTERS),
        ("WorkerTypes", R.WORKFORCE_TYPES),
        ("DirectionList", [d for d, _ in R.CAPACITY_DIRECTIONS]),
        ("LeavingReasons", R.LEAVING_REASONS),
        ("DispositionList", [d for d, _ in R.DISPOSITIONS]),
        ("LadderLevels", R.LADDER),
        ("CategoryList", [c for c, _, _ in R.RANK_CATEGORIES]),
        ("YesNo", ["Yes", "No"]),
    ] + list(extra)
    col = 1
    for name, values in lists:
        ws.cell(row=row, column=col, value=name.replace("List", "")).font = S.H2
        for i, v in enumerate(values, start=1):
            ws.cell(row=row + i, column=col, value=v)
        letter = get_column_letter(col)
        _name(wb, sheet, name, f"${letter}${row + 1}:${letter}${row + len(values)}")
        col += 2
    return col
